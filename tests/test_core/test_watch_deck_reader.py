"""The deck read cold. A client read the 2026-08-30 pack for the first time
and asked what "pairs", "cycles" and "read from: review site" meant, why a
brand had no coin packages, why the game lists were short, what the source
of the sentiment table was, and whether one brand really had more fairness
complaints than the others. These tests pin the answers into the deck."""

from __future__ import annotations

from pptx import Presentation

from core.watch_deck import _trends_facts, factual_narrative, render_executive_deck


def _card(n_brands: int = 3) -> dict:
    dims = [{"name": f"Dimension {i}", "weight_pct": 10.0} for i in range(4)]
    rows = []
    for i in range(n_brands):
        rows.append({
            "name": f"Brand {i}", "is_self": i == 0, "rank": None, "provisional": True,
            "overall": {"normalized_pct": 60.0 + i, "coverage_pct": 50.0},
            "dimensions": {d["name"]: {"score": 3} for d in dims[:2]},
        })
    return {"rows": rows, "dimensions": dims, "generated_at": "2026-09-02T06:00:00+00:00"}


def _render(tmp_path, card=None, **kw):
    card = card or _card()
    out = tmp_path / "d.pptx"
    render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                          gaps=[], evidence_count=4393, path=out, **kw)
    prs = Presentation(str(out))
    texts = []
    for sl in prs.slides:
        parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        for sh in sl.shapes:
            if sh.has_table:
                parts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
        texts.append("\n".join(parts))
    return prs, texts


def _voice(n_brands: int = 15) -> dict:
    themes = ["fairness_rtp", "promo_value", "redemption_speed", "payments"]
    field = {t: {"share": 0.2, "neg_share": 0.5} for t in themes}
    brands = []
    for i in range(n_brands):
        brands.append({
            "name": f"Brand {i}", "is_self": i == 0, "n": 74 if i < 12 else 3, "too_few": i >= 12,
            "neg_share": 0.5, "top_complaint": "fairness_rtp", "top_praise": "promo_value", "quotes": [],
            "themes": {t: {"share": 0.25, "neg_share": 0.6, "n": 18} for t in themes},
        })
    return {"window_days": 30, "mentions": 900, "sources": ["reddit", "app_store"], "min_mentions": 15,
            "field_themes": field, "brands": brands}


class TestReadingGuide:
    def test_slide_two_tells_a_stranger_how_to_read_the_deck(self, tmp_path) -> None:
        _prs, texts = _render(tmp_path, voice=_voice(3), catalog={
            "items": 5, "label": "", "totals": {"provider": 5}, "third_party_only": [],
            "brands": [{"subject_id": "b", "name": "Brand 1", "is_self": False, "counts": {"provider": 5},
                        "providers": ["A", "B"], "packages": [], "promotions": [], "games_sample": [], "sources": {}}]})
        guide = texts[2]                                                  # cover, contents, guide
        assert "HOW TO READ THIS DECK" in guide.upper()
        assert "3 brands" in guide and "4,393 facts" in guide and "4 dimensions" in guide
        assert "† Provisional" in guide and "rank is withheld" in guide
        assert "behind a login" in guide                                  # why a store can be empty
        assert "opinion" in guide and "Reddit" in guide                   # what "players say" is
        assert "raw inventory" in guide and "Click a name" in guide
        assert "Executive summary" in texts[3] or "EXECUTIVE SUMMARY" in texts[3].upper()   # after the guide

    def test_no_slide_says_pairs_or_cycles(self, tmp_path) -> None:
        _prs, texts = _render(tmp_path)
        joined = "\n".join(texts)
        assert "pairs" not in joined and "cycles" not in joined
        assert "brand × dimension cells" in joined


class TestTrendsReadPlainly:
    def test_one_mover_is_one_mover_and_late_starters_count_from_their_first_score(self) -> None:
        pts = [
            {"taken_at": "2026-07-25", "scores": {"A": None, "B": None}},
            {"taken_at": "2026-08-15", "scores": {"A": 53.0, "B": 60.0}},
            {"taken_at": "2026-08-27", "scores": {"A": 48.6, "B": None}},
        ]
        facts = _trends_facts({"points": pts, "brands": ["A", "B"]}, us=[])
        obs = " ".join(facts["observations"])
        assert "collection runs" in obs
        assert "Only one brand was scored at two runs – A: 53.0 → 48.6 (-4.4)" in obs
        assert "riser" not in obs and "faller" not in obs

    def test_largest_rise_and_fall_are_different_brands(self) -> None:
        pts = [
            {"taken_at": "2026-08-15", "scores": {"A": 50.0, "B": 60.0, "Us": 55.0}},
            {"taken_at": "2026-08-27", "scores": {"A": 58.0, "B": 52.0, "Us": 56.0}},
        ]
        facts = _trends_facts({"points": pts, "brands": ["A", "B", "Us"]}, us=["Us"])
        obs = " ".join(facts["observations"])
        assert "Largest rise – A: 50.0 → 58.0 (+8.0)" in obs
        assert "Largest fall – B: 60.0 → 52.0 (-8.0)" in obs
        assert "Us: 55.0 → 56.0 (+1.0) across 2 runs" in obs
        assert "flat lines" in facts["implications"][0]


class TestVoiceSlideReadsCold:
    def test_legend_sources_and_counts_are_visible_with_fifteen_brands(self, tmp_path) -> None:
        prs, texts = _render(tmp_path, voice=_voice(15))
        sl = next(sl for sl, t in zip(prs.slides, texts, strict=False) if "How to read: n = public posts" in t)
        text = texts[list(prs.slides).index(sl)]
        assert "Sources: Reddit (r/sweepstakescasinos), Apple App Store reviews" in text
        assert "grey = fewer than 15 posts" in text
        assert "25% (18)" in text                                          # share with its count
        tbl = next(sh for sh in sl.shapes if sh.has_table)
        assert tbl.top + tbl.height <= 6.2 * 914400                        # ends above the legend


class TestRawDataAnswersTheQuestions:
    def _catalog(self, packages, games):
        return {
            "items": 10, "label": "", "totals": {"provider": 2, "coin_package": len(packages), "promotion": 1,
                                                  "game": len(games)},
            "third_party_only": [],
            "brands": [{
                "subject_id": "b", "name": "Brand 1", "is_self": False,
                "counts": {"provider": 2, "coin_package": len(packages), "promotion": 1, "game": len(games)},
                "providers": ["Studio A", "Studio B"], "packages": packages,
                "promotions": [], "promotions_full": [
                    {"name": "Daily login bonus", "detail": "", "benefit": "1,500 GC", "how_to_claim": "Log in",
                     "frequency": "Daily", "url": "https://brand.example/promotions", "image": ""}],
                "games_sample": games[:12], "games_full": games, "sources": {},
                "customer_states": ["logged_out"], "observed_at": "2026-09-01",
            }],
        }

    def test_packages_have_a_description_and_link_to_their_page(self, tmp_path) -> None:
        pk = [{"name": "$9.99", "price_usd": 9.99, "coins": "", "detail": "110 VIP points; 1 Golden Key",
               "url": "https://brand.example/store", "source_type": "site", "gold_coins": 173500.0, "sweeps_coins": None}]
        prs, texts = _render(tmp_path, catalog=self._catalog(pk, [f"Game {i}" for i in range(30)]))
        page = next(t for t in texts if "Brand 1 – Coins / Promotions" in t)
        assert "Coin package | Gold coins | Sweeps coins | Description" in page
        assert "$9.99 | 173,500 | – | 110 VIP points; 1 Golden Key" in page
        assert "click a name to open the page" in page
        sl = next(sl for sl, t in zip(prs.slides, texts, strict=False) if "Brand 1 – Coins / Promotions" in t)
        links = {sh.table.cell(1, 0).text_frame.paragraphs[0].runs[0].hyperlink.address
                 for sh in sl.shapes if sh.has_table}
        assert links == {"https://brand.example/store", "https://brand.example/promotions"}

    def test_an_empty_store_and_a_short_game_list_say_why(self, tmp_path) -> None:
        _prs, texts = _render(tmp_path, catalog=self._catalog([], ["Money Train 2", "Lion Gems"]))
        coins = next(t for t in texts if "Brand 1 – Coins / Promotions" in t)
        assert "Coin store not read yet – it sits behind a login" in coins
        library = next(t for t in texts if "Brand 1 – Providers / Games" in t)
        assert "Only 2 titles appeared on the public pages" in library
        assert "has not been read as a signed-in player yet" in library

    def test_the_cross_brand_summaries_are_gone(self, tmp_path) -> None:
        _prs, texts = _render(tmp_path, catalog=self._catalog(
            [{"name": "$1.99", "price_usd": 1.99, "coins": "", "detail": "", "url": "u", "source_type": "site",
              "gold_coins": None, "sweeps_coins": None}], ["G1"]))
        joined = "\n".join(texts)
        assert "promotions on record – " not in joined.replace("packages and 1 promotions on record", "")
        assert "Game libraries observed" not in joined
        assert "Coin packages, as reported by public sources" not in joined


class TestNarratorRules:
    def test_the_narrator_is_told_to_write_for_a_first_time_reader(self) -> None:
        from tools.watch.tools import _DECK_NARRATIVE_SYSTEM as sysm

        assert "first time" in sysm and "game-format merchandising" in sysm
        assert "vs_field_pts" in sysm and "in line with the field" in sysm
        assert "draws more negativity" in sysm


class TestTrendsChartReadsAtAGlance:
    def test_axis_follows_the_data_empty_runs_go_and_movers_are_plotted(self, tmp_path) -> None:
        """2026-09-02: a 0–100 axis flattened a twenty-point band into one
        line, the first run was an empty column, and the panel named a
        faller the chart did not plot."""
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        card = _card(7)
        card["rows"][0]["overall"]["normalized_pct"] = 67.0          # Brand 0 is us
        trends = {
            "cycles": 3, "brands": [f"Brand {i}" for i in range(7)] + ["Ghost"],
            "points": [
                {"taken_at": "2026-07-25", "scores": {"Ghost": 40.0}},                        # nobody plotted
                {"taken_at": "2026-08-15", "scores": {f"Brand {i}": 60.0 + i for i in range(6)}},
                {"taken_at": "2026-08-27", "scores": {**{f"Brand {i}": 61.0 + i for i in range(5)}, "Brand 5": 52.0}},
            ],
        }
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, trends=trends)
        prs = Presentation(str(out))
        sl = next(sl for sl in prs.slides if any(sh.has_text_frame and "collection runs" in sh.text_frame.text for sh in sl.shapes))
        chart = next(sh for sh in sl.shapes if sh.has_chart).chart
        cats = list(chart.plots[0].categories)
        assert cats == ["2026-08-15", "2026-08-27"]                               # the empty run is gone
        assert chart.value_axis.minimum_scale == 40 and chart.value_axis.maximum_scale == 80   # 52…67, padded to tens
        names = [s.name for s in chart.series]
        assert "Brand 5" in names                                                  # the largest faller is plotted
        assert len({s.format.line.color.rgb for s in chart.series}) >= 5           # lines you can tell apart
        text = "\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
        assert "axis 40–80 of 100" in text


class TestLoyaltyPagesOnlyWithSubstance:
    def test_name_only_tiers_get_no_page(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        def brand(name, tiers):
            return {"subject_id": name, "name": name, "is_self": False,
                    "counts": {"provider": 1, "loyalty_tier": len(tiers)}, "providers": ["A"],
                    "packages": [], "promotions": [], "games_sample": [], "sources": {}, "tiers": tiers,
                    "customer_states": ["logged_out"], "observed_at": "2026-09-02"}
        catalog = {"items": 10, "label": "", "totals": {"provider": 2, "loyalty_tier": 9}, "third_party_only": [],
                   "brands": [brand("Brand F", [{"name": f"Tier {i}", "qualification": "", "reward": "", "url": "u"} for i in range(7)]),
                              brand("Brand A", [{"name": "Bronze", "qualification": "0–499 VIP points", "reward": "400,000 GC on $9.99", "url": "u"},
                                              {"name": "Ghost", "qualification": "", "reward": "", "url": "u"}])]}
        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=catalog)
        from pptx import Presentation
        texts = []
        for sl in Presentation(str(out)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
            texts.append("\n".join(parts))
        assert not any("Brand F – Loyalty Club" in t for t in texts)          # names only: no page
        page = next(t for t in texts if "Brand A – Loyalty Club" in t)
        assert "Bronze | 0–499 VIP points | 400,000 GC on $9.99" in page
        assert "Ghost" not in page                                                # an empty row is dropped
