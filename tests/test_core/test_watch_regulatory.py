"""Regulatory register (docs/88 §D): bills, effective dates, enforcement,
lawsuits, operator responses — third-party, verified excerpts, a calendar,
alerts, and on top of the pack only when present."""

from __future__ import annotations

import json
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace

import pytest
import pytest_asyncio

from core.watch import WatchManager
from core.watch_regulatory import (
    dedupe_key,
    extract_regulatory,
    regulatory_calendar,
    regulatory_queries,
    state_name,
)


def _d(days: int) -> str:
    return (datetime.now(UTC) + timedelta(days=days)).isoformat()[:10]


@pytest_asyncio.fixture
async def wm(tmp_path):
    from core.database import Database

    db = Database(str(tmp_path / "reg.db"))
    await db.initialize()
    yield WatchManager(db)
    await db.close()


class _Router:
    def __init__(self, items):
        self.items = items

    async def complete(self, **kw):
        return SimpleNamespace(content=json.dumps({"items": self.items}))


class TestExtraction:
    def test_queries_and_names(self) -> None:
        qs = regulatory_queries(["FL", "New York"], ["Brand C"], year=2026)
        assert any("Florida sweepstakes casino bill 2026" in q for q in qs)
        assert any('"Brand C" lawsuit' in q for q in qs)
        assert state_name("ny") == "New York" and state_name("Texas") == "Texas"

    @pytest.mark.asyncio
    async def test_items_need_a_verbatim_excerpt_valid_kind_and_jurisdiction(self) -> None:
        page = ("Governor Hochul signed S5935 on July 21, 2026, banning dual-currency sweepstakes casinos; "
                "the law takes effect on September 19, 2026. Separately, the Michigan Gaming Control Board "
                "issued cease-and-desist letters to nine operators including Brand C.")
        r = _Router([
            {"kind": "effective_date", "jurisdiction": "NY", "title": "S5935 sweepstakes ban takes effect", "status": "signed",
             "event_date": "2026-09-19", "subjects": [], "excerpt": "the law takes effect on September 19, 2026"},
            {"kind": "enforcement", "jurisdiction": "Michigan", "title": "MGCB C&D to nine operators", "status": "issued",
             "event_date": "", "subjects": ["Brand C", "Nope"], "excerpt": "issued cease-and-desist letters to nine operators"},
            {"kind": "rumor", "jurisdiction": "NY", "title": "x", "excerpt": "the law takes effect on September 19, 2026"},
            {"kind": "bill", "jurisdiction": "Mars", "title": "x", "excerpt": "the law takes effect on September 19, 2026"},
            {"kind": "bill", "jurisdiction": "CA", "title": "AB 831", "event_date": "soon", "excerpt": "AB 831 sailed through committee"},
        ])
        items = await extract_regulatory(r, page_text=page, brands=["Brand C"], states=["NY", "MI"])
        assert [(i["kind"], i["jurisdiction"]) for i in items] == [("effective_date", "NY"), ("enforcement", "MI")]
        assert items[1]["subjects"] == ["Brand C"]  # unknown brand dropped
        assert dedupe_key("ny", "bill", "S5935 ban", "2026-09-19") == dedupe_key("NY", "bill", "s5935 ban!", "2026-09-19")

    def test_calendar_reading(self) -> None:
        items = [
            {"jurisdiction": "NY", "kind": "effective_date", "title": "ban effective", "event_date": _d(30), "observed_at": _d(-1)},
            {"jurisdiction": "CA", "kind": "bill", "title": "AB 831 hearing", "event_date": _d(200), "observed_at": _d(-1)},
            {"jurisdiction": "MI", "kind": "enforcement", "title": "C&D", "event_date": "", "observed_at": _d(-2)},
            {"jurisdiction": "MI", "kind": "operator_response", "title": "Brand J exits Michigan", "event_date": _d(-10), "observed_at": _d(-5)},
        ]
        cal = regulatory_calendar(items, horizon_days=90)
        assert [i["jurisdiction"] for i in cal["ahead"]] == ["NY"]
        assert [i["title"] for i in cal["recent_actions"]] == ["C&D"]
        assert cal["operator_responses"][0]["title"] == "Brand J exits Michigan"
        assert cal["by_jurisdiction"]["MI"] == 2 and cal["total"] == 4


class TestManagerAndPack:
    @pytest.mark.asyncio
    async def test_add_list_calendar_alerts_report(self, wm) -> None:
        from tools.watch.tools import WatchAlertsTool, WatchBoardReportTool, WatchRegulatoryTool

        await wm.upsert_dimension(name="Promo", company_id="c1", weight_pct=100, subcriteria=[{"name": "w", "weight_pct": 100}])
        await wm.add_subject(company_id="c1", name="Us", is_self=True)
        r = await wm.add_regulatory(company_id="c1", jurisdiction="ny", kind="effective_date",
                                    title="S5935 sweepstakes ban takes effect", status="signed", event_date=_d(20),
                                    source_url="https://news/ny", excerpt="the law takes effect")
        assert r and r["jurisdiction"] == "NY"
        assert await wm.add_regulatory(company_id="c1", jurisdiction="NY", kind="effective_date",
                                       title="S5935 sweepstakes ban takes effect", event_date=_d(20)) is None
        with pytest.raises(ValueError):
            await wm.add_regulatory(company_id="c1", jurisdiction="NY", kind="gossip", title="x")
        await wm.add_regulatory(company_id="c1", jurisdiction="MI", kind="enforcement", title="MGCB C&D to nine operators",
                                subjects=["Brand C"], source_url="https://mgcb")
        cal = await wm.regulatory_calendar("c1")
        assert cal["total"] == 2 and cal["ahead"][0]["jurisdiction"] == "NY" and cal["recent_actions"][0]["jurisdiction"] == "MI"
        rt = WatchRegulatoryTool()
        rt._watch_manager = wm
        lst = await rt.execute({"company_id": "c1", "action": "list", "jurisdiction": "mi"})
        assert lst.data["count"] == 1
        # alerts: the effective date is inside 30 days, the enforcement is fresh
        at = WatchAlertsTool()
        at._watch_manager = wm
        at._gateway = None
        a = await at.execute({"company_id": "c1", "notify": False})
        kinds = sorted((x["kind"], x["subject_name"]) for x in a.data["new"])
        assert kinds == [("regulatory", "MI"), ("regulatory", "NY")]
        # report section + snapshot section
        rep = WatchBoardReportTool()
        rep._watch_manager, rep._router, rep._config = wm, None, None
        out = await rep.execute({"company_id": "c1", "baseline": True, "take_snapshot": True, "deck": False})
        md = out.data["markdown"]
        assert "## Regulatory calendar" in md and "| NY | effective date | S5935" in md and "MGCB C&D" in md
        snap = await wm.get_snapshot((await wm.latest_snapshot("c1"))[0])
        assert snap["regulatory"]["total"] == 2
        off = await rep.execute({"company_id": "c1", "baseline": True, "take_snapshot": False, "deck": False, "voice": "false"})
        assert "## Regulatory calendar" not in off.data["markdown"]

    @pytest.mark.asyncio
    async def test_collect_tool_searches_reads_and_files(self, wm, monkeypatch) -> None:
        import core.watch_observe as O
        from tools.watch import tools as T

        await wm.add_subject(company_id="c1", name="Brand C")

        async def fake_search(q, *, api_key, max_results=8, timeout=30.0):
            return [{"title": "NY signs ban", "url": "https://news.example/ny", "snippet": ""},
                    {"title": "dup", "url": "https://news.example/ny", "snippet": ""}]

        async def fake_fetch(url, *, browser_manager=None, proxy_url=None, timeout=20.0):
            return ("Governor signed S5935; the law takes effect on September 19, 2026.", None, "http")

        monkeypatch.setattr(O, "search_web", fake_search)
        monkeypatch.setattr(O, "fetch_page_best_effort", fake_fetch)
        t = T.WatchRegulatoryCollectTool()
        t._watch_manager = wm
        t._router = _Router([{"kind": "effective_date", "jurisdiction": "NY", "title": "S5935 sweepstakes ban takes effect",
                              "status": "signed", "event_date": "2026-09-19", "subjects": [],
                              "excerpt": "the law takes effect on September 19, 2026"}])
        t._vault = {"search_sh_api_key": "k"}
        t._config = None
        r = await t.execute({"company_id": "c1", "states": ["NY"]})
        assert r.success, r.error
        assert r.data["urls"] == 1 and r.data["pages_read"] == 1 and r.data["filed"] == 1
        assert r.data["items"][0]["source_url"] == "https://news.example/ny"
        r2 = await t.execute({"company_id": "c1", "states": ["NY"]})
        assert r2.data["filed"] == 0 and r2.data["duplicates"] == 1
        nokey = T.WatchRegulatoryCollectTool()
        nokey._watch_manager, nokey._router, nokey._vault = wm, t._router, {}
        assert not (await nokey.execute({"company_id": "c1"})).success


class TestDeck:
    def test_calendar_slide_only_when_items_exist(self, tmp_path) -> None:
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [{"name": "Us", "is_self": True, "rank": 1, "provisional": False,
                          "overall": {"normalized_pct": 60.0, "coverage_pct": 70.0}, "dimensions": {}}], "dimensions": []}
        cal = regulatory_calendar([
            {"jurisdiction": "NY", "kind": "effective_date", "title": "S5935 sweepstakes ban takes effect", "status": "signed",
             "event_date": _d(20), "observed_at": _d(-1), "source_url": "u"},
            {"jurisdiction": "MI", "kind": "enforcement", "title": "MGCB C&D to nine operators", "event_date": "", "observed_at": _d(-1)},
        ])
        a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
        n = factual_narrative(card, None, [], [])
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=a)
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=b, regulatory=cal)
        assert len(Presentation(str(b)).slides) == len(Presentation(str(a)).slides) + 1
        texts = []
        for sl in Presentation(str(b)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
            texts.append("\n".join(parts))
        slide = next(t for t in texts if "REGULATORY CALENDAR" in t.upper())
        assert "S5935" in slide and "MGCB C&D" in slide and "not legal advice" in slide
