"""Tone of voice — how each brand talks to players.

Built from copy already in the registers; the habits are measured, the
quotes are verbatim, and a characterisation the model did not back with a
real line is dropped."""

from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from pptx import Presentation

from core.watch_tone import copy_samples, read_tone, summarize_tone, tone_features


class _Router:
    def __init__(self, payload):
        self.payload = payload
        self.seen = None

    async def complete(self, messages, **kw):
        self.seen = messages[-1]["content"]
        return SimpleNamespace(content=json.dumps(self.payload))


class TestMeasurements:
    def test_shouting_urgency_and_length_are_measured(self) -> None:
        loud = tone_features(["GET YOUR 150% EXTRA COINS NOW! Limited time only!"])
        calm = tone_features(["Play our games and redeem prizes when you are ready."])
        assert loud["caps_pct"] > calm["caps_pct"]
        assert loud["urgency_per_100w"] > calm["urgency_per_100w"]
        assert loud["exclaims_per_line"] == 2.0 and calm["exclaims_per_line"] == 0.0
        assert tone_features([]) == {"samples": 0}

    def test_emoji_and_second_person(self) -> None:
        f = tone_features(["🎰 Your daily bonus is ready, and your free spins are waiting"])
        assert f["emoji_per_line"] == 1.0 and f["second_person_per_100w"] > 0


class TestSamples:
    def test_copy_comes_from_promotions_emails_and_site_lines_deduped(self) -> None:
        catalog = [
            SimpleNamespace(kind="promotion", name="150% Extra Coins", detail="First purchase only",
                            source_url="https://b/promo", observed_at="2026-08-27T00:00:00"),
            SimpleNamespace(kind="game", name="Sweet Bonanza", detail="", source_url="", observed_at=""),
        ]
        comms = [SimpleNamespace(subject_line="Your bonus is waiting", offer_text="150% Extra Coins",
                                 received_at="2026-08-26T09:00:00")]
        evidence = [
            {"dimension": "Marketing proposition", "value_text": "America's #1 social casino",
             "excerpt": "", "source_url": "https://b", "observed_at": "2026-08-25"},
            {"dimension": "Payment options", "value_text": "Visa accepted", "excerpt": "",
             "source_url": "", "observed_at": ""},
        ]
        got = copy_samples(brand="B", evidence=evidence, catalog=catalog, comms=comms)
        texts = [g["text"] for g in got]
        assert "150% Extra Coins" in texts and "Your bonus is waiting" in texts
        assert "America's #1 social casino" in texts
        assert "Sweet Bonanza" not in texts  # a game title is not copy
        assert "Visa accepted" not in texts  # payments is not marketing
        assert len(texts) == len(set(texts))  # the promo said twice, kept once
        assert {g["source"] for g in got} >= {"promotion", "e-mail subject", "site copy"}


class TestReading:
    @pytest.mark.asyncio
    async def test_a_signature_the_brand_never_said_is_dropped(self) -> None:
        brands = [{
            "name": "B", "is_self": False,
            "samples": [{"text": "GET YOUR 150% EXTRA COINS", "source": "promotion"}],
            "features": tone_features(["GET YOUR 150% EXTRA COINS"]),
        }]
        r = _Router({"brands": [
            {"brand": "B", "register": "loud and urgent", "traits": ["shouts in capitals"],
             "signature": "Something it never wrote", "avoid": ""},
            {"brand": "Ghost", "register": "invented", "traits": [], "signature": ""},
        ]})
        got = await read_tone(r, brands)
        assert set(got) == {"B"}  # a brand with no samples is never characterised
        assert got["B"]["register"] == "loud and urgent" and got["B"]["signature"] == ""

    @pytest.mark.asyncio
    async def test_no_router_leaves_the_measurements_speaking(self) -> None:
        brands = [{"name": "B", "is_self": True, "samples": [{"text": "Play free today"}],
                   "features": tone_features(["Play free today"])}]
        assert await read_tone(None, brands) == {}
        s = summarize_tone(brands, {})
        assert s["source"] == "measurements" and s["brands"][0]["register"] == ""

    def test_summary_ranks_loudest_and_most_urgent_ours_first(self) -> None:
        brands = [
            {"name": "Peer", "is_self": False, "samples": [{"text": "HURRY! LAST CHANCE NOW!"}],
             "features": tone_features(["HURRY! LAST CHANCE NOW!"])},
            {"name": "Us", "is_self": True, "samples": [{"text": "Play a few free games when you like"}],
             "features": tone_features(["Play a few free games when you like"])},
        ]
        s = summarize_tone(brands)
        assert s["brands"][0]["name"] == "Us"
        assert s["loudest"] == "Peer" and s["most_urgent"] == "Peer"


class TestSlide:
    def test_slide_appears_only_with_copy(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [{"name": "Peer", "is_self": False, "rank": 1, "provisional": False,
                          "overall": {"normalized_pct": 60.0, "coverage_pct": 70.0}, "dimensions": {}}],
                "dimensions": []}
        tone = summarize_tone(
            [{"name": "Peer", "is_self": False,
              "samples": [{"text": "GET YOUR 150% EXTRA COINS NOW!", "source": "promotion"}],
              "features": tone_features(["GET YOUR 150% EXTRA COINS NOW!"])}],
            {"Peer": {"register": "loud and urgent", "traits": ["shouts"],
                      "signature": "GET YOUR 150% EXTRA COINS NOW!", "avoid": "no emoji"}},
        )
        a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
        n = factual_narrative(card, None, [], [])
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=a)
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=b, tone=tone)
        assert len(Presentation(str(b)).slides) == len(Presentation(str(a)).slides) + 1
        texts = []
        for sl in Presentation(str(b)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
            texts.append("\n".join(parts))
        slide = next(t for t in texts if "TONE OF VOICE" in t.upper())
        assert "loud and urgent" in slide and "GET YOUR 150% EXTRA COINS NOW!" in slide
