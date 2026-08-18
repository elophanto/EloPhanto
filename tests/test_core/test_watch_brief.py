"""The weekly service (docs/88 §A–B): the Friday brief and mid-week alerts,
read from registers that already exist."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta

import pytest
import pytest_asyncio
from pptx import Presentation

from core.watch import WatchManager
from core.watch_alerts import alerts_from_events, alerts_from_regulatory, alerts_from_voice
from core.watch_brief import (
    brief_facts,
    build_weekly_brief,
    field_changes,
    offer_changes,
    render_brief_markdown,
    render_brief_slide,
)


def _iso(days_ago: float) -> str:
    return (datetime.now(UTC) - timedelta(days=days_ago)).isoformat()


class TestFieldChanges:
    def test_newest_in_window_vs_newest_before_and_new_coverage(self) -> None:
        since = _iso(7)
        ev = [  # newest first
            {"subject": "Crown", "dimension": "Promotions", "subcriterion": "welcome", "claim": "200% extra", "value_text": "200%", "observed_at": _iso(1), "source_url": "u1"},
            {"subject": "Crown", "dimension": "Promotions", "subcriterion": "welcome", "claim": "150% extra", "value_text": "150%", "observed_at": _iso(3), "source_url": "u1b"},
            {"subject": "Crown", "dimension": "Promotions", "subcriterion": "welcome", "claim": "150% extra", "value_text": "150%", "observed_at": _iso(20), "source_url": "u0"},
            {"subject": "Crown", "dimension": "Loyalty", "subcriterion": "tiers", "claim": "5 tiers", "value_text": "5", "observed_at": _iso(2), "source_url": "u2"},
            {"subject": "Crown", "dimension": "Loyalty", "subcriterion": "tiers", "claim": "5 tiers", "value_text": "5", "observed_at": _iso(30), "source_url": "u2"},
            {"subject": "Spree", "dimension": "Packages", "subcriterion": "min", "claim": "$4.99", "value_text": "4.99", "observed_at": _iso(1), "source_url": "u3"},
        ]
        got = field_changes(ev, since=since)
        by = {(c["brand"], c["dimension"]): c for c in got}
        assert by[("Crown", "Promotions")]["before"] == "150%" and by[("Crown", "Promotions")]["after"] == "200%"
        assert not by[("Crown", "Promotions")]["new_coverage"]
        assert ("Crown", "Loyalty") not in by  # unchanged
        assert by[("Spree", "Packages")]["new_coverage"] and by[("Spree", "Packages")]["before"] is None

    def test_offer_changes_before_after(self) -> None:
        since = _iso(7)
        card = {"rows": [{"name": "Crown", "is_self": False, "rank": 1, "overall": {"normalized_pct": 60.0}}]}
        ev = [
            {"subject": "Crown", "dimension": "Promotional proposition", "claim": "New players get 200% extra Gold Coins on first purchase.", "value_text": "", "observed_at": _iso(1), "source_url": "u1"},
            {"subject": "Crown", "dimension": "Promotional proposition", "claim": "New players get 150% extra Gold Coins on first purchase.", "value_text": "", "observed_at": _iso(20), "source_url": "u0"},
        ]
        got = offer_changes(card, ev, since=since)
        assert len(got) == 1 and got[0]["field"] == "welcome" and "150%" in got[0]["before"] and "200%" in got[0]["after"]


@pytest_asyncio.fixture
async def wm(tmp_path):
    from core.database import Database

    db = Database(str(tmp_path / "brief.db"))
    await db.initialize()
    yield WatchManager(db)
    await db.close()


class TestBrief:
    @pytest.mark.asyncio
    async def test_build_render_and_tool(self, wm, tmp_path) -> None:
        from tools.watch.tools import WatchWeeklyBriefTool

        await wm.upsert_dimension(name="Promotional proposition and generosity", company_id="c1", weight_pct=60,
                                  subcriteria=[{"name": "welcome", "weight_pct": 100}])
        await wm.upsert_dimension(name="State availability and variation", company_id="c1", weight_pct=40,
                                  subcriteria=[{"name": "states", "weight_pct": 100}])
        us = await wm.add_subject(company_id="c1", name="Us", is_self=True)
        cc = await wm.add_subject(company_id="c1", name="Crown")
        ll = await wm.add_subject(company_id="c1", name="LuckyLand")
        dims = {d.name: d.dimension_id for d in await wm.list_dimensions("c1")}
        promo, states = dims["Promotional proposition and generosity"], dims["State availability and variation"]
        # last month: 150%; this week: 200%
        await wm.add_evidence(company_id="c1", subject_id=cc.subject_id, dimension_id=promo, subcriterion="welcome",
                              claim="New players get 150% extra Gold Coins on first purchase.", value_text="150%",
                              observed_at=_iso(20), source_url="https://c/old")
        await wm.add_evidence(company_id="c1", subject_id=cc.subject_id, dimension_id=promo, subcriterion="welcome",
                              claim="New players get 200% extra Gold Coins on first purchase.", value_text="200%",
                              observed_at=_iso(1), source_url="https://c/new")
        await wm.add_evidence(company_id="c1", subject_id=us.subject_id, dimension_id=promo, subcriterion="welcome",
                              claim="Welcome Bundle with 100% extra Gold Coins.", value_text="100%",
                              observed_at=_iso(2), source_url="https://us")
        await wm.add_evidence(company_id="c1", subject_id=ll.subject_id, dimension_id=states, subcriterion="states",
                              claim="LuckyLand Slots is closing on September 14, 2026.", value_text="2026-09-14",
                              observed_at=_iso(1), source_url="https://ll")
        brief = await build_weekly_brief(wm, "c1", request="Is Crown's welcome offer bigger than ours?",
                                         answer="Yes — 200% vs our 100%.")
        assert brief["new_facts"] == 3 and brief["brands_touched"] == 3
        real = [c for c in brief["field_changes"] if not c["new_coverage"]]
        assert len(real) == 1 and real[0]["brand"] == "Crown" and real[0]["after"] == "200%"
        assert brief["offer_changes"] and brief["offer_changes"][0]["brand"] == "Us"  # ours listed first
        assert brief["market_events"][0]["brand"] == "LuckyLand"
        facts = brief_facts(brief)
        assert "1 field change" in facts["changed"][0]
        assert any("LuckyLand" in m for m in facts["market"])
        assert any("Decide the response to LuckyLand" in d for d in facts["decisions"])
        md = render_brief_markdown(brief, {})
        assert "# Weekly competitive brief" in md and "| Crown | Promotional proposition and generosity · welcome | 150% | 200% |" in md
        assert "Is Crown's welcome offer bigger" in md
        p = render_brief_slide(brief, {"headline": "Crown raised its welcome offer to 200%"}, path=tmp_path / "b.pptx")
        prs = Presentation(p)
        assert len(prs.slides) == 1
        txt = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
        assert "Crown raised its welcome offer to 200%" in txt and "DECISIONS" in txt.upper()
        # the tool: writes both, snapshots, reports notified=False without a gateway
        t = WatchWeeklyBriefTool()
        t._watch_manager = wm
        t._router = None
        t._gateway = None
        r = await t.execute({"company_id": "c1", "path": str(tmp_path / "w" / "brief.md"), "notify": True})
        assert r.success and set(r.data["written"]) == {"markdown", "slide"} and r.data["notified"] is False
        assert r.data["snapshot_id"] and r.data["narrative_source"] == "facts"


class TestAlerts:
    def test_candidates_from_events_voice_and_regulatory(self) -> None:
        ev = alerts_from_events([{"brand": "LuckyLand", "claim": "LuckyLand Slots is closing on September 14, 2026.", "url": "u"}])
        assert ev[0]["kind"] == "market_event" and ev[0]["dedupe_key"]
        voice = {"since": _iso(7), "window_days": 7, "brands": [
            {"name": "Crown", "too_few": False, "themes": {"redemption_speed": {"n": 12, "share": 0.5, "neg_share": 0.8}}},
            {"name": "Spin", "too_few": False, "themes": {"support": {"n": 12, "share": 0.5, "neg_share": 0.8}}},
        ]}
        vdiff = {"baseline": False, "changed": [
            {"brand": "Crown", "theme": "redemption_speed", "neg_from": 0.4, "neg_to": 0.8, "direction": "rising"}]}
        vs = alerts_from_voice(voice, vdiff)
        assert [a["subject_name"] for a in vs] == ["Crown"]  # Spin: no rise recorded
        first = alerts_from_voice(voice, {"baseline": True, "changed": []})
        assert {a["subject_name"] for a in first} == {"Crown", "Spin"}
        soon = (datetime.now(UTC) + timedelta(days=10)).isoformat()[:10]
        far = (datetime.now(UTC) + timedelta(days=100)).isoformat()[:10]
        reg = alerts_from_regulatory([
            {"jurisdiction": "NY", "kind": "effective_date", "title": "SB 5935 sweepstakes ban", "event_date": soon, "source_url": "r1"},
            {"jurisdiction": "CA", "kind": "effective_date", "title": "AB 831", "event_date": far},
            {"jurisdiction": "MI", "kind": "enforcement", "title": "C&D to 9 operators", "observed_at": _iso(0.5)},
            {"jurisdiction": "MI", "kind": "enforcement", "title": "old C&D", "observed_at": _iso(10)},
        ])
        assert [(a["subject_name"], a["kind"]) for a in reg] == [("NY", "regulatory"), ("MI", "regulatory")]

    @pytest.mark.asyncio
    async def test_check_stores_once_and_reports_notify_state(self, wm) -> None:
        from tools.watch.tools import WatchAlertsTool

        await wm.upsert_dimension(name="State availability and variation", company_id="c1", weight_pct=100,
                                  subcriteria=[{"name": "states", "weight_pct": 100}])
        ll = await wm.add_subject(company_id="c1", name="LuckyLand")
        dim = (await wm.list_dimensions("c1"))[0]
        await wm.add_evidence(company_id="c1", subject_id=ll.subject_id, dimension_id=dim.dimension_id,
                              subcriterion="states", claim="LuckyLand Slots is closing on September 14, 2026.",
                              observed_at=_iso(0.2), source_url="https://ll")

        class _GW:
            def __init__(self):
                self.sent = []

            async def broadcast(self, msg, session_id=None):
                self.sent.append(msg)

        t = WatchAlertsTool()
        t._watch_manager = wm
        gw = _GW()
        t._gateway = gw
        r1 = await t.execute({"company_id": "c1"})
        assert r1.success and len(r1.data["new"]) == 1 and r1.data["notified"] is True
        assert gw.sent and gw.sent[0].data["notification_type"] == "watch" and "LuckyLand" in gw.sent[0].data["text"]
        r2 = await t.execute({"company_id": "c1"})
        assert r2.data["new"] == [] and r2.data["notified"] is False  # stored once
        lst = await t.execute({"company_id": "c1", "action": "list"})
        assert lst.data["count"] == 1 and lst.data["alerts"][0]["notified_at"]
