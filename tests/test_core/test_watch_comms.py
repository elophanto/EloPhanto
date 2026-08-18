"""Player comms (docs/88 §C): what brands send players, from the organ's
own inboxes — read into watch_comms, shown on top of the pack, never a
score."""

from __future__ import annotations

import json
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace

import pytest
import pytest_asyncio

from core.watch import WatchManager
from core.watch_comms import (
    COMMS_CATEGORIES,
    excerpt_is_verbatim,
    inbox_username,
    read_comms,
)


def _iso(days_ago: float) -> str:
    return (datetime.now(UTC) - timedelta(days=days_ago)).isoformat()


@pytest_asyncio.fixture
async def wm(tmp_path):
    from core.database import Database

    db = Database(str(tmp_path / "comms.db"))
    await db.initialize()
    yield WatchManager(db)
    await db.close()


class _Router:
    def __init__(self, items):
        self.items = items

    async def complete(self, **kw):
        return SimpleNamespace(content=json.dumps({"items": self.items}))


class TestReading:
    def test_inbox_names_do_not_advertise_and_excerpts_must_be_verbatim(self) -> None:
        name = inbox_username("Crown Coins Casino")
        assert name.startswith("crown-coins-casino-") and len(name.split("-")[-1]) == 4
        assert excerpt_is_verbatim("200% extra Gold Coins", "Get 200% EXTRA gold coins today!")
        assert not excerpt_is_verbatim("double coins", "Get 200% extra gold coins today!")
        assert "promo_offer" in COMMS_CATEGORIES

    @pytest.mark.asyncio
    async def test_read_validates_category_and_keeps_row_when_excerpt_fails(self) -> None:
        msgs = [
            {"id": "m1", "subject": "Your 200% welcome boost", "text": "Get 200% extra Gold Coins on your first purchase today only."},
            {"id": "m2", "subject": "We miss you", "text": "Come back for 5,000 free coins."},
        ]
        r = _Router([
            {"id": "m1", "category": "promo_offer", "offer": "200% extra GC on first purchase", "excerpt": "200% extra Gold Coins on your first purchase"},
            {"id": "m2", "category": "reactivation", "offer": "5,000 free coins", "excerpt": "we want you back"},  # paraphrase
        ])
        items, dropped = await read_comms(r, brand="Crown", messages=msgs)
        assert [i["category"] for i in items] == ["promo_offer", "reactivation"]
        assert items[0]["excerpt"] and items[1]["excerpt"] == ""  # kept, quote dropped
        assert dropped["unverified_excerpt"] == 1


class TestManagerAndSummary:
    @pytest.mark.asyncio
    async def test_add_dedupes_by_message_id_and_summary_reads_cadence(self, wm) -> None:
        cc = await wm.add_subject(company_id="c1", name="Crown")
        await wm.tag_subject(cc.subject_id, "inbox:crown-ab12@agentmail.to")
        for i in range(6):
            r = await wm.add_comms(company_id="c1", subject_id=cc.subject_id, message_id=f"<m{i}@x>",
                                   received_at=_iso(i * 2 + 0.5), category="promo_offer" if i % 2 else "daily_bonus",
                                   inbox="crown-ab12@agentmail.to", sender="Crown <no-reply@crown>",
                                   subject_line=f"Offer {i}", offer_text=f"{100 + i}% extra" if i % 2 else "", excerpt="x")
            assert r is not None and r.weekday is not None
        assert await wm.add_comms(company_id="c1", subject_id=cc.subject_id, message_id="<m0@x>",
                                  received_at=_iso(1), category="other") is None
        with pytest.raises(ValueError):
            await wm.add_comms(company_id="c1", subject_id=cc.subject_id, message_id="<z>", received_at=_iso(1), category="spam")
        s = await wm.comms_summary("c1", window_days=14)
        b = s["brands"][0]
        assert s["emails"] == 6 and b["n"] == 6 and b["per_week"] == 3.0
        assert b["categories"] == {"daily_bonus": 3, "promo_offer": 3}
        assert b["latest_offers"][0]["offer"] == "101% extra"
        # the pack is untouched
        assert await wm.list_evidence("c1") == [] and await wm.list_scores("c1") == []
        snap = await wm.get_snapshot(await wm.take_snapshot("c1"))
        assert snap["comms"]["emails"] == 6


class _FakeAgentMail:
    def __init__(self):
        me = self
        self.created = []
        self.mail = {
            "crown-ab12@agentmail.to": [
                SimpleNamespace(message_id="<a@crown>", subject="200% welcome boost", from_="Crown <hi@crown>",
                                timestamp=_iso(1), preview="Get 200% extra Gold Coins", extracted_text="Get 200% extra Gold Coins on your first purchase."),
                SimpleNamespace(message_id="<b@crown>", subject="Daily wheel is live", from_="Crown <hi@crown>",
                                timestamp=_iso(0.5), preview="Spin the daily wheel", extracted_text="Spin the daily wheel for free coins every day."),
            ]
        }

        class _Messages:
            def list(self, inbox_id):
                return SimpleNamespace(messages=me.mail.get(inbox_id, []))

            def get(self, inbox_id, message_id):
                return next(m for m in me.mail.get(inbox_id, []) if m.message_id == message_id)

        class _Inboxes:
            messages = _Messages()

            def create(self, request):
                addr = f"{request.username}@agentmail.to"
                me.created.append(addr)
                me.mail.setdefault(addr, [])
                return SimpleNamespace(inbox_id=addr)

        self.inboxes = _Inboxes()


class TestTools:
    @pytest.mark.asyncio
    async def test_setup_then_collect_then_pack_picks_it_up(self, wm, tmp_path, monkeypatch) -> None:
        import sys
        import types

        from tools.watch import tools as T

        fake = _FakeAgentMail()
        mod = types.ModuleType("agentmail")
        mod.AgentMail = lambda api_key: fake
        sub = types.ModuleType("agentmail.inboxes.types")
        sub.CreateInboxRequest = lambda **kw: SimpleNamespace(**kw)
        monkeypatch.setitem(sys.modules, "agentmail", mod)
        monkeypatch.setitem(sys.modules, "agentmail.inboxes.types", sub)
        vault = {"agentmail_api_key": "k"}

        class _V(dict):
            def set(self, k, v):
                self[k] = v

        vault = _V(vault)
        cfg = SimpleNamespace(email=SimpleNamespace(api_key_ref="agentmail_api_key"))
        await wm.add_subject(company_id="c1", name="Us", is_self=True)
        cc = await wm.add_subject(company_id="c1", name="Crown", url="https://crown.example")
        # register is canon
        setup = T.WatchCommsSetupTool()
        setup._watch_manager, setup._vault, setup._config = wm, vault, cfg
        bad = await setup.execute({"subject": "Fortune Coins", "company_id": "c1"})
        assert not bad.success
        # pre-link Crown to the fake inbox that has mail (setup is idempotent on a linked inbox)
        await wm.tag_subject(cc.subject_id, "inbox:crown-ab12@agentmail.to")
        ok = await setup.execute({"subject": "Crown", "company_id": "c1"})
        assert ok.success and ok.data["inbox"] == "crown-ab12@agentmail.to" and "sign up" in ok.data["next"]
        assert vault[f"watch_comms:{cc.subject_id}"]["email"] == "crown-ab12@agentmail.to"
        # a fresh brand gets a created inbox
        us = await wm.get_subject_by_name("Us", "c1")
        ok2 = await setup.execute({"subject": "Us", "company_id": "c1"})
        assert ok2.success and fake.created and ok2.data["inbox"].startswith("us-")
        assert f"inbox:{ok2.data['inbox']}" in (await wm.get_subject_by_name("Us", "c1")).tags
        del us
        # collect
        col = T.WatchCommsCollectTool()
        col._watch_manager, col._vault, col._config = wm, vault, cfg
        col._router = _Router([
            {"id": "<a@crown>", "category": "promo_offer", "offer": "200% extra GC on first purchase", "excerpt": "200% extra Gold Coins on your first purchase"},
            {"id": "<b@crown>", "category": "daily_bonus", "offer": "", "excerpt": "Spin the daily wheel for free coins"},
        ])
        r = await col.execute({"company_id": "c1"})
        assert r.success and r.data["kept_total"] == 2
        crown = next(b for b in r.data["brands"] if b["subject"] == "Crown")
        assert crown["fetched"] == 2 and crown["kept"] == 2
        r2 = await col.execute({"company_id": "c1"})
        assert r2.data["kept_total"] == 0  # message ids dedupe
        # read side + pack
        rd = T.WatchCommsTool()
        rd._watch_manager = wm
        summ = await rd.execute({"company_id": "c1"})
        assert summ.data["emails"] == 2
        rep = T.WatchBoardReportTool()
        rep._watch_manager, rep._router, rep._config = wm, None, None
        await wm.upsert_dimension(name="Promo", company_id="c1", weight_pct=100, subcriteria=[{"name": "w", "weight_pct": 100}])
        out = await rep.execute({"company_id": "c1", "baseline": True, "take_snapshot": False, "deck": False})
        assert "## What they send players" in out.data["markdown"] and "| Crown | 2 |" in out.data["markdown"]
        off = await rep.execute({"company_id": "c1", "baseline": True, "take_snapshot": False, "deck": False, "voice": "false"})
        assert "## What they send players" not in off.data["markdown"]


class TestDeck:
    def test_comms_slide_only_when_present(self, tmp_path) -> None:
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [{"name": "Us", "is_self": True, "rank": 1, "provisional": False,
                          "overall": {"normalized_pct": 60.0, "coverage_pct": 70.0}, "dimensions": {}}], "dimensions": []}
        comms = {"window_days": 30, "since": "", "emails": 5, "label": "What brands send players — x",
                 "brands": [{"subject_id": "s", "name": "Us", "is_self": True, "inbox": "us@agentmail.to", "n": 5, "per_week": 1.2,
                             "categories": {"promo_offer": 3, "daily_bonus": 2}, "peak_hour_utc": 15,
                             "latest_offers": [{"received_at": "2026-08-17", "category": "promo_offer", "offer": "200% extra", "subject": "x"}],
                             "latest_subjects": ["x"]}]}
        a = tmp_path / "a.pptx"
        b = tmp_path / "b.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []), gaps=[], evidence_count=1, path=a)
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], [], comms=comms), gaps=[],
                              evidence_count=1, path=b, comms=comms)
        na, nb = len(Presentation(str(a)).slides), len(Presentation(str(b)).slides)
        assert nb == na + 1
        texts = []
        for sl in Presentation(str(b)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
            texts.append("\n".join(parts))
        slide = next(t for t in texts if "What they send players" in t)
        assert "Us  (us)" in slide and "1.2" in slide and "200% extra" in slide
