"""docs/88 §E — message categories, demand calendar, app meta, trends."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta

import pytest
import pytest_asyncio
from pptx import Presentation

from core.watch import WatchManager
from core.watch_calendar import demand_calendar, message_categories, message_category, us_holidays


class TestCategoriesAndCalendar:
    def test_message_categories_are_a_fixed_vocabulary(self) -> None:
        assert message_category("New players can claim a Welcome Bundle with 200% extra") == "welcome_offer"
        assert message_category("Get 150% more coins on your first purchase") == "purchase_bonus"
        assert message_category("Spin the daily wheel") == "daily_login"
        assert message_category("Halloween slot race with $10k jackpot drop") == "tournament_race"
        assert message_category("Terms apply") == "other"
        ev = [
            {"subject": "Crown", "dimension": "Promotional proposition and generosity", "claim": "Welcome bundle 200%", "observed_at": "2026-08-17"},
            {"subject": "Crown", "dimension": "Promotional proposition and generosity", "claim": "Daily login bonus", "observed_at": "2026-08-17"},
            {"subject": "Crown", "dimension": "Game portfolio", "claim": "New slots weekly", "observed_at": "2026-08-17"},
            {"subject": "Spree", "dimension": "Loyalty programme", "claim": "VIP tiers with points", "observed_at": "2026-07-01"},
        ]
        mc = message_categories(ev, since="2026-08-01")
        assert mc["brands"] == {"Crown": {"welcome_offer": 1, "daily_login": 1}}
        assert message_categories(ev)["field"]["vip_loyalty"] == 1

    def test_demand_calendar_is_computed_and_events_are_supplied(self) -> None:
        cal = demand_calendar(weeks=8, now="2026-08-18T00:00:00+00:00",
                              events=[{"date": "2026-09-10", "label": "NFL kickoff", "kind": "sports"}])
        labels = [i["label"] for i in cal["items"]]
        assert "Labor Day" in labels and "SSI payment" in labels and "Payday (15th)" in labels
        assert "NFL kickoff" in labels and cal["weeks"][0]["week_of"] == "2026-08-18"
        assert cal["items"][0]["date"] >= "2026-08-18" and cal["items"][-1]["date"] <= cal["to"]
        assert dict(us_holidays(2026))[datetime(2026, 11, 26).date()] == "Thanksgiving"
        # no sports unless supplied
        assert not any(i["kind"] == "sports" for i in demand_calendar(weeks=8, now="2026-08-18T00:00:00+00:00")["items"])


@pytest_asyncio.fixture
async def wm(tmp_path):
    from core.database import Database

    db = Database(str(tmp_path / "smalls.db"))
    await db.initialize()
    yield WatchManager(db)
    await db.close()


class TestAppMetaAndTrends:
    @pytest.mark.asyncio
    async def test_app_meta_history_flags_a_new_release_in_the_brief(self, wm) -> None:
        from core.watch_brief import brief_facts, build_weekly_brief

        await wm.upsert_dimension(name="Promo", company_id="c1", weight_pct=100, subcriteria=[{"name": "w", "weight_pct": 100}])
        cc = await wm.add_subject(company_id="c1", name="Crown")
        await wm.add_app_meta(company_id="c1", subject_id=cc.subject_id, store="app_store", app_id="1", version="3.1", rating=4.5, rating_count=100)
        await wm.add_app_meta(company_id="c1", subject_id=cc.subject_id, store="app_store", app_id="1", version="3.2", rating=4.4,
                              rating_count=120, release_notes="Bug fixes and a new lobby")
        latest = await wm.app_meta_latest("c1")
        m = latest[cc.subject_id]
        assert m["version"] == "3.2" and m["previous_version"] == "3.1"
        brief = await build_weekly_brief(wm, "c1")
        assert brief["releases"] == [{"brand": "Crown", "version": "3.2", "was": "3.1", "notes": "Bug fixes and a new lobby", "rating": 4.4}]
        assert any("shipped app v3.2 (was 3.1)" in c for c in brief_facts(brief)["changed"])
        assert brief["calendar_next"]  # the next two weeks always have paydays / benefit dates

    @pytest.mark.asyncio
    async def test_trends_need_three_scored_cycles_and_render(self, wm, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        await wm.upsert_dimension(name="Promo", company_id="c1", weight_pct=100, subcriteria=[{"name": "w", "weight_pct": 100}])
        us = await wm.add_subject(company_id="c1", name="Us", is_self=True)
        dim = (await wm.list_dimensions("c1"))[0]
        await wm.add_evidence(company_id="c1", subject_id=us.subject_id, dimension_id=dim.dimension_id, subcriterion="w",
                              claim="c", source_url="u")
        for score in (2.0, 3.0, 4.0):
            await wm.set_score(company_id="c1", subject_id=us.subject_id, dimension_id=dim.dimension_id, score=score,
                               rationale="x", subcriteria_scores={"w": score})
            await wm.take_snapshot("c1", label=f"cycle {score}")
        tr = await wm.trend_series("c1")
        # three snapshots on the same day collapse to one point — the pack shows trends only across days
        assert tr["cycles"] == 1
        # simulate three days by rewriting taken_at
        rows = await wm._db.execute("SELECT snapshot_id FROM watch_snapshots WHERE company_id='c1' ORDER BY taken_at", ())
        for i, r in enumerate(rows):
            day = (datetime.now(UTC) - timedelta(days=(3 - i) * 30)).isoformat()
            await wm._db.execute("UPDATE watch_snapshots SET taken_at=? WHERE snapshot_id=?", (day, r["snapshot_id"]))
        tr = await wm.trend_series("c1")
        assert tr["cycles"] == 3 and [pt["scores"]["Us"] for pt in tr["points"]] == [40.0, 60.0, 80.0]
        card = await wm.scorecard("c1")
        a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
        n = factual_narrative(card, None, [], [])
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=a)
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=b, trends=tr,
                              calendar=demand_calendar(weeks=8))
        assert len(Presentation(str(b)).slides) == len(Presentation(str(a)).slides) + 2
        texts = ["\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in Presentation(str(b)).slides]
        assert any("Scores over 3 cycles" in t for t in texts) and any("Demand calendar" in t.title() or "DEMAND CALENDAR" in t for t in texts)
