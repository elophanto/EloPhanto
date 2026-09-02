"""Voice of customer in the deliverables (docs/87, step 4): on top of the
pack when present, absent otherwise, and on its own."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta

import pytest
import pytest_asyncio
from pptx import Presentation

from core.watch import WatchManager
from core.watch_deck import factual_narrative, render_executive_deck, render_voice_deck
from core.watch_xlsx import render_scorecard_xlsx, render_voice_xlsx


def _card():
    return {
        "rows": [
            {"name": "Us", "is_self": True, "rank": 1, "provisional": False,
             "overall": {"normalized_pct": 60.0, "coverage_pct": 70.0}, "dimensions": {}},
            {"name": "Crown", "is_self": False, "rank": 2, "provisional": False,
             "overall": {"normalized_pct": 55.0, "coverage_pct": 65.0}, "dimensions": {}},
        ],
        "dimensions": [],
    }


def _voice(*, crown_n=20):
    themes = {"redemption_speed": {"n": crown_n // 2, "share": 0.5, "neg_share": 0.9},
              "game_selection": {"n": crown_n // 2, "share": 0.5, "neg_share": 0.1}}
    return {
        "window_days": 30, "since": "", "mentions": crown_n + 3, "min_mentions": 15,
        "sources": ["reddit", "app_store"],
        "label": "What players say — sentiment from public posts, not observed product fact.",
        "field_themes": themes,
        "brands": [
            {"subject_id": "us", "name": "Us", "is_self": True, "n": 3, "too_few": True, "sources": ["reddit"],
             "avg_rating": None, "neg_share": 0.0, "pos_share": 1.0, "themes": {}, "top_complaint": None,
             "top_praise": None, "quotes": [], "flags": []},
            {"subject_id": "cc", "name": "Crown", "is_self": False, "n": crown_n, "too_few": False,
             "sources": ["reddit", "app_store"], "avg_rating": 3.4, "neg_share": 0.5, "pos_share": 0.4,
             "themes": themes, "top_complaint": "redemption_speed", "top_praise": "game_selection",
             "quotes": [{"quote": "Redemption took nine days", "source": "reddit", "url": "u", "posted_at": "2026-08-10",
                         "theme": "redemption_speed", "sentiment": "negative"}],
             "flags": ["KYC strategy and customer journey"]},
        ],
    }


def _texts(path):
    prs = Presentation(str(path))
    out = []
    for sl in prs.slides:
        parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        for sh in sl.shapes:
            if sh.has_table:
                parts += [c.text for r in sh.table.rows for c in r.cells]
        out.append("\n".join(parts))
    return out


class TestPackUnchangedWithoutVoice:
    def test_no_voice_means_the_deck_is_exactly_the_pack(self, tmp_path) -> None:
        card = _card()
        a = tmp_path / "a.pptx"
        b = tmp_path / "b.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=3, path=a)
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=3, path=b, voice=None, voice_diff=None)
        assert len(_texts(a)) == len(_texts(b))
        assert not any("What players say" in t for t in _texts(b))
        # an empty summary is treated as absent too
        c = tmp_path / "c.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=3, path=c, voice={"mentions": 0, "brands": []})
        assert len(_texts(c)) == len(_texts(a))


class TestVoiceOnTop:
    def test_voice_adds_the_slides_and_the_profile_strip(self, tmp_path) -> None:
        card = _card()
        voice = _voice()
        narrative = factual_narrative(card, None, [], [], voice=voice)
        assert narrative["slides"]["voice"]["observations"]
        narrative["profiles"] = [{"brand": "Crown", "title": "Crown", "observations": ["o1", "o2"], "implications": ["i1"]}]
        out = tmp_path / "v.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=narrative, gaps=[], evidence_count=3,
                              path=out, voice=voice, voice_diff={"baseline": True, "changed": [], "material_count": 0})
        texts = _texts(out)
        heat = next(t for t in texts if "What players say about the field" in t)
        assert "Crown" in heat and "too few (n<" in heat          # short, so the row stays one line
        assert "not observed product fact" in heat
        assert any("Sentiment baseline set" in t for t in texts)
        profile = next(t for t in texts if "COMPETITOR DEEP DIVE" in t.upper() and "\nCrown\n" in t)
        assert "WHAT PLAYERS SAY · N=20" in profile and "Redemption took nine days" in profile
        assert "players flag KYC strategy and customer journey" in profile

    def test_movement_slide_lists_rising_and_falling(self, tmp_path) -> None:
        card = _card()
        voice = _voice()
        vdiff = {"baseline": False, "material_count": 1, "changed": [
            {"brand": "Crown", "is_self": False, "theme": "redemption_speed", "n_from": 5, "n_to": 10,
             "share_from": 0.2, "share_to": 0.5, "neg_from": 0.5, "neg_to": 0.9, "direction": "rising"}]}
        out = tmp_path / "m.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], [], voice=voice),
                              gaps=[], evidence_count=3, path=out, voice=voice, voice_diff=vdiff)
        mv = next(t for t in _texts(out) if "1 theme moved this cycle" in t)
        assert "▲ Crown – Redemption speed: 20% → 50% of mentions, 50% → 90% negative (n 5 → 10)" in mv


class TestStandalone:
    def test_voice_deck_and_workbook_on_their_own(self, tmp_path) -> None:
        voice = _voice()
        out = tmp_path / "voice.pptx"
        render_voice_deck(voice, voice_diff=None, narrative=factual_narrative({"rows": [], "dimensions": []}, None, [], [], voice=voice), path=out)
        texts = _texts(out)
        assert any("What players say about the field" in t for t in texts)
        assert any("Crown – players complain about redemption speed" in t for t in texts)
        assert any("How to read these numbers" in t for t in texts)
        # brands too few to read get no slide of their own
        assert not any("Us – " in t for t in texts)
        rows = [{"brand": "Crown", "source": "reddit", "posted_at": "2026-08-10", "theme": "redemption_speed",
                 "sentiment": "negative", "rating": None, "quote": "Redemption took nine days", "dimension": "KYC",
                 "geo_hint": "", "weight": 0.85, "url": "u"}]
        xp = render_voice_xlsx(rows, path=tmp_path / "voice.xlsx", summary=voice)
        from openpyxl import load_workbook

        wb = load_workbook(xp)
        assert wb.sheetnames == ["Summary", "Voice"]
        assert wb["Voice"]["G5"].value == "Redemption took nine days"

    def test_scorecard_workbook_gains_a_voice_sheet_only_when_rows_given(self, tmp_path) -> None:
        card = _card()
        from openpyxl import load_workbook

        a = render_scorecard_xlsx(card, dimensions=[], evidence=[], staleness=[], path=tmp_path / "a.xlsx")
        assert "Voice" not in load_workbook(a).sheetnames
        b = render_scorecard_xlsx(card, dimensions=[], evidence=[], staleness=[], path=tmp_path / "b.xlsx",
                                  voice_rows=[{"brand": "Crown", "source": "reddit", "quote": "q", "theme": "support",
                                               "sentiment": "negative"}])
        assert "Voice" in load_workbook(b).sheetnames


@pytest_asyncio.fixture
async def wm(tmp_path):
    from core.database import Database

    db = Database(str(tmp_path / "vd.db"))
    await db.initialize()
    yield WatchManager(db)
    await db.close()


class TestReportTools:
    @pytest.mark.asyncio
    async def test_board_report_has_no_voice_section_until_rows_exist_then_auto_adds_it(self, wm, tmp_path) -> None:
        from tools.watch.tools import WatchBoardReportTool, WatchVoiceReportTool

        await wm.upsert_dimension(name="KYC strategy and customer journey", company_id="c1", weight_pct=100,
                                  subcriteria=[{"name": "kyc", "weight_pct": 100}])
        us = await wm.add_subject(company_id="c1", name="Us", is_self=True)
        cc = await wm.add_subject(company_id="c1", name="Crown")
        dim = (await wm.list_dimensions("c1"))[0]
        await wm.add_evidence(company_id="c1", subject_id=us.subject_id, dimension_id=dim.dimension_id,
                              subcriterion="kyc", claim="ID check at first redemption", source_url="https://us")
        await wm.set_score(company_id="c1", subject_id=us.subject_id, dimension_id=dim.dimension_id,
                           score=4.0, rationale="x", subcriteria_scores={"kyc": 4.0})
        t = WatchBoardReportTool()
        t._watch_manager = wm
        t._router = None
        t._config = None
        r0 = await t.execute({"company_id": "c1", "baseline": True, "take_snapshot": False, "deck": False})
        assert r0.success and "## What players say" not in r0.data["markdown"]
        # standalone refuses with nothing collected
        vt = WatchVoiceReportTool()
        vt._watch_manager = wm
        v0 = await vt.execute({"company_id": "c1", "take_snapshot": False})
        assert not v0.success and "watch_voice_collect" in v0.error
        when = (datetime.now(UTC) - timedelta(days=2)).isoformat()
        for i in range(16):
            await wm.add_voice(company_id="c1", subject_id=cc.subject_id, source="reddit", theme="redemption_speed",
                               sentiment="negative", quote=f"redemption took forever number {i}", posted_at=when,
                               source_url=f"https://r/{i}")
        r1 = await t.execute({"company_id": "c1", "baseline": True, "take_snapshot": False, "deck": False})
        md = r1.data["markdown"]
        assert "## What players say" in md and "| Crown | 16 | 100% | redemption speed |" in md
        assert "too few mentions to read" in md  # Us
        r2 = await t.execute({"company_id": "c1", "baseline": True, "take_snapshot": False, "deck": False, "voice": "false"})
        assert "## What players say" not in r2.data["markdown"]
        # standalone pack
        out = tmp_path / "voice" / "voice.md"
        v1 = await vt.execute({"company_id": "c1", "path": str(out), "take_snapshot": True})
        assert v1.success, v1.error
        assert set(v1.data["written"]) == {"report", "deck", "workbook"} and not v1.data["errors"]
        assert v1.data["snapshot_id"] and v1.data["brands_readable"] == 1
