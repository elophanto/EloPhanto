"""How old is the page? — the watch organ reads the date a page was
written, not only the day it read it, and a third-party claim past the
horizon (or one the brand's own signed-in lobby did not confirm) is shown
as *reported*, never as *carried*.

2026-09-02: a studio that had left one of our brands was still marked as
carried, on the strength of a review written before it left."""

from __future__ import annotations

from datetime import date

import pytest


class TestPageDate:
    def test_structured_markup_wins_and_modified_beats_published(self) -> None:
        from core.watch_observe import page_date

        html = (
            '<html><head><meta property="article:published_time" content="2024-03-01T10:00:00Z">'
            '<meta property="article:modified_time" content="2026-07-15T08:30:00+00:00">'
            "</head><body>Last updated: January 1, 2020</body></html>"
        )
        assert page_date(html) == ("2026-07-15", "high")

    def test_attribute_order_and_json_ld_and_time_tags(self) -> None:
        from core.watch_observe import page_date

        assert page_date('<meta content="2026-05-02" name="date">') == (
            "2026-05-02",
            "high",
        )
        assert page_date(
            '<script type="application/ld+json">{"datePublished":"2025-11-30T00:00:00"}</script>'
        ) == ("2025-11-30", "high")
        assert page_date('<time datetime="2026-01-09T12:00:00Z">9 Jan</time>') == (
            "2026-01-09",
            "high",
        )

    def test_visible_text_is_medium_and_many_forms_parse(self) -> None:
        from core.watch_observe import page_date, parse_date_text

        assert page_date(
            "<p>Review last updated on September 2, 2026 by our team</p>"
        ) == ("2026-09-02", "medium")
        assert page_date("<p>Published: 2 September 2026</p>") == (
            "2026-09-02",
            "medium",
        )
        assert page_date("<p>Updated 09/02/2026</p>") == ("2026-09-02", "medium")
        assert parse_date_text("Sept. 2, 2026") == "2026-09-02"
        assert parse_date_text("nonsense") == ""

    def test_header_is_low_and_nothing_is_nothing(self) -> None:
        from core.watch_observe import page_date

        assert page_date(
            "<p>no dates here</p>", {"Last-Modified": "Wed, 02 Sep 2026 10:00:00 GMT"}
        ) == ("2026-09-02", "low")
        assert page_date("<p>no dates here</p>") == ("", "")

    def test_future_and_ancient_dates_are_not_believed(self) -> None:
        from core.watch_observe import page_date

        assert page_date('<meta name="date" content="2099-01-01">') == ("", "")
        assert page_date('<meta name="date" content="1999-01-01">') == ("", "")


class TestSearchEngineDates:
    def test_since_and_freshness_boost_are_sent_only_when_given(self) -> None:
        from core.watch_observe import search_payload

        plain = search_payload("q")
        assert "since" not in plain and "freshness_boost" not in plain
        body = search_payload("q", since="2026-03-01", freshness_boost=True)
        assert body["since"] == "2026-03-01" and body["freshness_boost"] is True

    def test_a_header_only_modified_date_is_not_the_best_date(self) -> None:
        """Search.sh: a modified_at read from HTTP Last-Modified is the
        site's deploy time — an upper bound, never an update date."""
        from core.watch_observe import source_dates

        marked_up = source_dates({"published_at": "2024-02-01", "modified_at": "2026-06-30T14:12:00.000Z",
                                  "date_confidence": "high",
                                  "date_sources": {"modified_at": "json-ld:dateModified"}})
        assert marked_up["best_date"] == "2026-06-30"
        header = source_dates({"published_at": "2024-02-01", "modified_at": "2026-09-01",
                               "date_confidence": "low", "date_sources": {"modified_at": "http:last-modified"}})
        assert header["best_date"] == "2024-02-01" and header["modified_at"] == "2026-09-01"
        nothing = source_dates({"published_at": None, "modified_at": None, "date_confidence": None})
        assert nothing == {"published_at": "", "modified_at": "", "date_confidence": "", "best_date": ""}

    def test_the_engines_conflicts_come_back_normalised(self) -> None:
        from core.watch_observe import _search_conflicts

        got = _search_conflicts({"conflicts": [
            {"summary": "a source from 2026-07 says the studio left the brand; an older one still lists it",
             "dated": True,
             "newer": {"claim": "the studio left the brand in June 2026", "source_url": "https://n.example/",
                       "source_title": "n", "date": "2026-07-14T08:00:00.000Z"},
             "older": {"claim": "the studio is part of the brand", "source_url": "https://o.example/",
                       "source_title": "o", "date": "2024-03-02"}},
            "not a dict",
        ]})
        assert len(got) == 1
        c = got[0]
        assert c["dated"] and c["newer"]["date"] == "2026-07-14" and c["older"]["date"] == "2024-03-02"
        assert c["newer"]["url"] == "https://n.example/" and "left the brand" in c["newer"]["claim"]
        assert _search_conflicts({"sources": []}) == [] and _search_conflicts(None) == []

    def test_sources_carry_dates_when_the_engine_gives_them(self) -> None:
        from core.watch_observe import _search_sources

        got = _search_sources(
            {
                "sources": [
                    {
                        "url": "https://a.example/x",
                        "title": "A",
                        "snippet": "s",
                        "published_at": "2024-02-01T00:00:00Z",
                        "modified_at": "2026-08-01",
                        "date_confidence": "high",
                    },
                    {"url": "https://b.example/y", "title": "B", "snippet": ""},
                    {"url": "ftp://nope", "title": "no"},
                ]
            }
        )
        assert [g["url"] for g in got] == ["https://a.example/x", "https://b.example/y"]
        assert (
            got[0]["published_at"] == "2024-02-01"
            and got[0]["modified_at"] == "2026-08-01"
        )
        assert got[0]["date_confidence"] == "high"
        assert got[1]["published_at"] == "" and got[1]["modified_at"] == ""

    @pytest.mark.asyncio
    async def test_an_engine_that_refuses_since_is_asked_again_without_it(
        self, monkeypatch
    ) -> None:
        import core.watch_observe as wo

        calls: list[dict] = []

        class _Resp:
            def __init__(self, code, data):
                self.status_code, self._data, self.text = code, data, ""

            def json(self):
                return self._data

        class _Client:
            def __init__(self, **kw):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *a):
                return False

            async def post(self, url, *, json, headers):
                calls.append(json)
                if "since" in json:
                    return _Resp(422, {"error": "unknown field since"})
                return _Resp(
                    200,
                    {
                        "sources": [
                            {"url": "https://r.example/p", "title": "t", "snippet": ""}
                        ]
                    },
                )

        import httpx

        monkeypatch.setattr(httpx, "AsyncClient", _Client)
        got = await wo.search_web("q", api_key="k", since="2026-03-01")
        assert [g["url"] for g in got] == ["https://r.example/p"]
        assert "since" in calls[0] and "since" not in calls[1]

    @pytest.mark.asyncio
    async def test_the_dated_search_returns_sources_conflicts_and_answer(self, monkeypatch) -> None:
        import core.watch_observe as wo

        class _Resp:
            status_code = 200
            text = ""

            def json(self):
                return {
                    "answer": "The studio left in June 2026.",
                    "sources": [{"url": "https://n.example/", "title": "n", "snippet": "",
                                 "published_at": "2026-07-14", "modified_at": None, "date_confidence": "high"}],
                    "conflicts": [{"summary": "newer says left, older says carried", "dated": True,
                                   "newer": {"claim": "left", "source_url": "https://n.example/", "date": "2026-07-14"},
                                   "older": {"claim": "carried", "source_url": "https://o.example/", "date": "2024-03-02"}}],
                }

        class _Client:
            def __init__(self, **kw):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *a):
                return False

            async def post(self, url, *, json, headers):
                assert json["since"] == "2026-03-01" and json["freshness_boost"] is True
                return _Resp()

        import httpx
        monkeypatch.setattr(httpx, "AsyncClient", _Client)
        got = await wo.search_web_dated("q", api_key="k", since="2026-03-01", freshness_boost=True)
        assert got["sources"][0]["best_date"] == "2026-07-14"
        assert got["conflicts"][0]["older"]["date"] == "2024-03-02"
        assert got["answer"].startswith("The studio left")

    def test_dated_old_results_sink_below_undated_ones(self) -> None:
        from core.watch_catalog import rank_research_urls

        got = rank_research_urls(
            [{"url": "https://old.example/brand", "modified_at": "2024-01-01"},
             {"url": "https://undated.example/brand"},
             {"url": "https://fresh.example/brand", "published_at": "2026-08-20"},
             # a deploy-time modified date does not rescue an old page
             {"url": "https://deployed.example/brand", "published_at": "2023-05-05", "modified_at": "2026-09-01",
              "best_date": "2023-05-05"}],
            today=date(2026, 9, 2),
            limit=10,
        )
        assert got.index("https://old.example/brand") > got.index("https://undated.example/brand")
        assert got.index("https://deployed.example/brand") > got.index("https://undated.example/brand")
    def test_provider_and_game_queries_carry_the_year(self) -> None:
        from core.watch_catalog import research_queries

        qs = [
            q for _, q in research_queries("Brand K", ["provider", "game"], year=2026)
        ]
        assert any("2026" in q for q in qs[:2]) and any("2026" in q for q in qs[2:])


class TestHorizon:
    def test_age_and_staleness(self) -> None:
        from core.watch_catalog import STALE_AFTER_DAYS, is_stale, page_age_days

        today = date(2026, 9, 2)
        assert page_age_days("2026-08-02", today) == 31
        assert (
            page_age_days("", today) is None and page_age_days("garbage", today) is None
        )
        assert not is_stale("2026-08-02", today)
        assert is_stale("2025-12-01", today)
        assert not is_stale("", today)  # unknown is not stale
        assert STALE_AFTER_DAYS == 180


class TestMatrixKnowsWhatIsCurrent:
    """The cell says what the evidence supports today."""

    BRANDS = [
        {"name": "Brand A", "is_self": True},
        {"name": "Brand J", "is_self": False},
    ]
    TODAY = date(2026, 9, 2)

    def _cell(self, items, brand="Brand A", studio="pragmatic"):
        from core.watch_catalog import provider_matrix

        m = provider_matrix(items, self.BRANDS, today=self.TODAY)
        row = next(r for r in m["providers"] if r["key"].startswith(studio[:6]))
        return m, row["brands"][brand]

    def test_a_dated_review_is_reported_not_carried(self) -> None:
        m, cell = self._cell(
            [
                {
                    "brand": "Brand A",
                    "kind": "provider",
                    "name": "Pragmatic Play",
                    "detail": "",
                    "source_type": "third_party",
                    "page_date": "2024-11-03",
                },
            ]
        )
        assert cell["reported"] and not cell["carried"]
        assert cell["reason"] == "page dated 2024-11-03"
        assert m["counts"]["reported_only"] == 1 and m["counts"]["observed"] == 0

    def test_a_recent_review_still_counts(self) -> None:
        _m, cell = self._cell(
            [
                {
                    "brand": "Brand A",
                    "kind": "provider",
                    "name": "Pragmatic Play",
                    "detail": "",
                    "source_type": "third_party",
                    "page_date": "2026-07-01",
                },
            ]
        )
        assert (
            cell["carried"]
            and not cell["reported"]
            and cell["page_date"] == "2026-07-01"
        )

    def test_an_undated_review_is_not_discounted(self) -> None:
        _m, cell = self._cell(
            [
                {
                    "brand": "Brand A",
                    "kind": "provider",
                    "name": "Pragmatic Play",
                    "detail": "",
                    "source_type": "third_party",
                },
            ]
        )
        assert cell["carried"]

    def test_the_signed_in_lobby_overrides_the_web(self) -> None:
        """The lobby read as a player lists Hacksaw and not Pragmatic; a
        fresh review lists Pragmatic. The lobby wins."""
        items = [
            {
                "brand": "Brand A",
                "kind": "provider",
                "name": "Hacksaw Gaming",
                "detail": "",
                "source_type": "site",
                "customer_state": "registered",
            },
            {
                "brand": "Brand A",
                "kind": "provider",
                "name": "Pragmatic Play",
                "detail": "",
                "source_type": "third_party",
                "page_date": "2026-08-20",
            },
        ]
        _m, prag = self._cell(items)
        assert prag["reported"] and prag["reason"] == "not in the signed-in lobby read"
        _m, hack = self._cell(items, studio="hacksaw")
        assert hack["carried"] and hack["source"] == "site"

    def test_a_logged_out_site_page_does_not_override_a_review(self) -> None:
        """The public page may be a teaser; only a signed-in lobby read is
        complete enough to contradict."""
        items = [
            {
                "brand": "Brand A",
                "kind": "provider",
                "name": "Hacksaw Gaming",
                "detail": "",
                "source_type": "site",
                "customer_state": "logged_out",
            },
            {
                "brand": "Brand A",
                "kind": "provider",
                "name": "Pragmatic Play",
                "detail": "",
                "source_type": "third_party",
                "page_date": "2026-08-20",
            },
        ]
        _m, prag = self._cell(items)
        assert prag["carried"]

    def test_the_brands_own_page_is_never_stale(self) -> None:
        _m, cell = self._cell(
            [
                {
                    "brand": "Brand A",
                    "kind": "provider",
                    "name": "Pragmatic Play",
                    "detail": "",
                    "source_type": "site",
                    "page_date": "2020-01-01",
                },
            ]
        )
        assert cell["carried"]

    def test_the_freshest_of_two_reviews_is_kept(self) -> None:
        _m, cell = self._cell(
            [
                {
                    "brand": "Brand A",
                    "kind": "provider",
                    "name": "Pragmatic Play",
                    "detail": "",
                    "source_type": "third_party",
                    "page_date": "2024-01-01",
                },
                {
                    "brand": "Brand A",
                    "kind": "provider",
                    "name": "Pragmatic",
                    "detail": "",
                    "source_type": "third_party",
                    "page_date": "2026-08-01",
                },
            ]
        )
        assert cell["carried"] and cell["page_date"] == "2026-08-01"

    def test_the_summary_names_what_was_reported_but_not_counted(self) -> None:
        from core.watch_catalog import summarize_catalog

        class Row:
            def __init__(
                self,
                brand,
                name,
                source_type="third_party",
                state="logged_out",
                meta=None,
            ):
                self.subject_id, self.kind, self.name, self.detail = (
                    brand,
                    "provider",
                    name,
                    "",
                )
                self.source_type, self.customer_state, self.observed_at = (
                    source_type,
                    state,
                    "2026-09-02",
                )
                self.price_usd, self.coins_text, self.sort_index, self.meta = (
                    None,
                    "",
                    0,
                    meta or {},
                )
                self.source_url, self.image_path = "u", ""

        class Subj:
            def __init__(self, sid, name, is_self=False):
                self.subject_id, self.name, self.is_self = sid, name, is_self

        cat = summarize_catalog(
            [
                Row("a", "Hacksaw Gaming", "site", "registered"),
                Row("a", "Pragmatic Play", meta={"page_date": "2024-05-05"}),
            ],
            [Subj("a", "Brand A", True)],
        )
        brand = cat["brands"][0]
        assert brand["providers_reported"] == [
            "Pragmatic Play (not in the signed-in lobby read)"
        ]
        assert cat["matrix"]["counts"]["reported_only"] == 1


class TestDeckAndWorkbookShowReported:
    def _catalog(self):
        return {
            "items": 3,
            "label": "",
            "totals": {"provider": 3},
            "third_party_only": [],
            "brands": [
                {
                    "subject_id": "a",
                    "name": "Brand A",
                    "is_self": True,
                    "counts": {"provider": 2},
                    "providers": ["Hacksaw Gaming", "Pragmatic Play"],
                    "providers_reported": ["Pragmatic Play (page dated 2024-05-05)"],
                    "packages": [],
                    "promotions": [],
                    "games_sample": [],
                    "sources": {},
                }
            ],
            "matrix": {
                "providers": [
                    {
                        "key": "pragmaticplay",
                        "name": "Pragmatic Play",
                        "on_client_list": True,
                        "brands": {
                            "Brand A": {
                                "carried": False,
                                "reported": True,
                                "source": "third_party",
                                "games": 0,
                                "page_date": "2024-05-05",
                                "reason": "page dated 2024-05-05",
                            }
                        },
                        "brand_count": 0,
                        "reported_count": 1,
                        "observed": False,
                    },
                    {
                        "key": "hacksawgaming",
                        "name": "Hacksaw Gaming",
                        "on_client_list": True,
                        "brands": {
                            "Brand A": {
                                "carried": True,
                                "reported": False,
                                "source": "site",
                                "games": 3,
                                "page_date": "",
                                "reason": "",
                            }
                        },
                        "brand_count": 1,
                        "reported_count": 0,
                        "observed": True,
                    },
                ],
                "brands": ["Brand A"],
                "counts": {
                    "observed": 1,
                    "on_client_list": 2,
                    "both": 1,
                    "list_only": 1,
                    "observed_only": 0,
                    "reported_only": 1,
                },
                "brands_only_on_client_list": [],
                "brands_only_in_register": [],
            },
        }

    def test_the_deck_marks_reported_cells_and_names_them(self, tmp_path) -> None:
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        card = {
            "rows": [
                {
                    "name": "Brand A",
                    "is_self": True,
                    "rank": 1,
                    "provisional": False,
                    "overall": {"normalized_pct": 50.0, "coverage_pct": 50.0},
                    "dimensions": {},
                }
            ],
            "dimensions": [],
            "generated_at": "2026-09-02T00:00:00+00:00",
        }
        out = tmp_path / "d.pptx"
        render_executive_deck(
            card,
            diff=None,
            judged=[],
            summary=factual_narrative(card, None, [], []),
            gaps=[],
            evidence_count=1,
            path=out,
            catalog=self._catalog(),
        )
        texts = [
            "\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
            for sl in Presentation(str(out)).slides
        ]
        tables = [
            sh.table
            for sl in Presentation(str(out)).slides
            for sh in sl.shapes
            if sh.has_table
        ]
        cells = {
            t.cell(r, c).text
            for t in tables
            for r in range(len(t.rows))
            for c in range(len(t.columns))
        }
        assert "○" in cells and "● 3" in cells
        portfolio = next(t for t in texts if "Game portfolio" in t)
        assert "○ reported only" in portfolio
        library = next(t for t in texts if "Providers / Games" in t)
        assert (
            "Reported but not counted" in library and "page dated 2024-05-05" in library
        )

    def test_the_workbook_marks_reported_cells(self, tmp_path) -> None:
        from openpyxl import Workbook, load_workbook

        from core.watch_xlsx import write_portfolio_sheet

        wb = Workbook()
        rows = [
            {
                "brand": "Brand A",
                "is_self": True,
                "kind": "provider",
                "name": "Hacksaw Gaming",
                "detail": "",
                "source_type": "site",
                "customer_state": "registered",
            },
            {
                "brand": "Brand A",
                "is_self": True,
                "kind": "provider",
                "name": "Pragmatic Play",
                "detail": "",
                "source_type": "third_party",
                "page_date": "2024-05-05",
            },
        ]
        write_portfolio_sheet(wb, rows)
        p = tmp_path / "w.xlsx"
        wb.save(p)
        ws = load_workbook(p)["Game portfolio"]
        marks = {
            ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(4, ws.max_row + 1)
        }
        assert marks["Hacksaw Gaming"] == "●" and marks["Pragmatic Play"] == "○"
        assert "○ reported only" in str(ws["A1"].value)
