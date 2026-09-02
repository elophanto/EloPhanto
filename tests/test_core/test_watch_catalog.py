"""Catalog — the raw inventory behind the scores (docs/89).

The client asked for "game providers, coin packages, promotions and game
list, just the raw data". It is a list, not a judgement: verified against
the page, stamped with the session it was read in, never scored.
"""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
import pytest_asyncio
from pptx import Presentation

from core.watch import WatchManager
from core.watch_catalog import (
    catalog_page_kind,
    dedupe_key,
    extract_catalog,
    parse_price,
    rank_catalog_pages,
)


class _Router:
    def __init__(self, items):
        self.items = items

    async def complete(self, **kw):
        return SimpleNamespace(content=json.dumps({"items": self.items}))


@pytest_asyncio.fixture
async def wm(tmp_path):
    from core.database import Database

    db = Database(str(tmp_path / "cat.db"))
    await db.initialize()
    yield WatchManager(db)
    await db.close()


class TestPagePicking:
    def test_each_kind_finds_its_page_and_legal_pages_are_never_one(self) -> None:
        assert catalog_page_kind("https://b.com/store", "Buy coins") == "coin_package"
        assert catalog_page_kind("https://b.com/promotions", "") == "promotion"
        assert catalog_page_kind("https://b.com/providers", "") == "provider"
        assert catalog_page_kind("https://b.com/slots", "All slots") == "game"
        assert catalog_page_kind("https://b.com/terms", "Terms") == ""
        assert catalog_page_kind("https://b.com/sweepstakes-rules", "") == ""

    def test_pages_group_by_kind_best_first(self) -> None:
        pages = [
            {"url": "https://b.com/", "title": "Home"},
            {"url": "https://b.com/store", "title": "Store"},
            {"url": "https://b.com/shop", "title": "Shop"},
            {"url": "https://b.com/buy", "title": "Buy"},
            {"url": "https://b.com/promotions", "title": "Promotions"},
        ]
        got = rank_catalog_pages(pages, per_kind=2)
        assert [p["url"] for p in got["coin_package"]] == ["https://b.com/store", "https://b.com/shop"]
        assert [p["url"] for p in got["promotion"]] == ["https://b.com/promotions"]

    def test_price_parsing(self) -> None:
        assert parse_price("$29.99 package") == 29.99
        assert parse_price("$4,999.00") == 4999.0
        assert parse_price("Best value") is None


class TestExtraction:
    @pytest.mark.asyncio
    async def test_items_must_be_printed_on_the_page(self) -> None:
        page = "Our studios: Pragmatic Play, Hacksaw Gaming and NetEnt power the lobby."
        r = _Router([
            {"name": "Pragmatic Play", "detail": "", "index": 0},
            {"name": "NetEnt", "detail": "", "index": 2},
            {"name": "Evolution Gaming", "detail": "invented", "index": 3},  # not on the page
        ])
        got = await extract_catalog(r, kind="provider", brand="B", page_text=page)
        assert [g["name"] for g in got] == ["Pragmatic Play", "NetEnt"]

    @pytest.mark.asyncio
    async def test_packages_keep_order_price_and_grant(self) -> None:
        page = "$9.99 gets GC 100,000 + free SC 10. $29.99 gets GC 700,000 + free SC 55."
        r = _Router([
            {"name": "$9.99", "coins": "GC 100,000 + free SC 10", "detail": "starter", "index": 0},
            {"name": "$29.99", "coins": "GC 700,000 + free SC 55", "detail": "", "index": 1},
        ])
        got = await extract_catalog(r, kind="coin_package", brand="B", page_text=page)
        assert [g["price_usd"] for g in got] == [9.99, 29.99]
        assert got[0]["coins_text"] == "GC 100,000 + free SC 10"
        assert [g["sort_index"] for g in got] == [0, 1]

    @pytest.mark.asyncio
    async def test_no_router_or_bad_kind_invents_nothing(self) -> None:
        assert await extract_catalog(None, kind="game", brand="B", page_text="x") == []
        assert await extract_catalog(_Router([]), kind="nonsense", brand="B", page_text="x") == []


class TestResearchFirst:
    """Public pages and the open web first; a session is spent only on what
    they cannot answer, and the metered exit is not spent at all."""

    def test_queries_name_the_brand_and_the_kind(self) -> None:
        from core.watch_catalog import research_queries

        qs = dict(research_queries("McLuck", ["provider", "coin_package"], year=2026))
        assert any("providers" in q for k, q in research_queries("McLuck", ["provider"], year=2026))
        assert "coin packages" in " ".join(
            q for k, q in research_queries("McLuck", ["coin_package"], year=2026)
        )
        assert qs  # both kinds produced queries

    def test_the_brands_own_domain_outranks_coupon_farms(self) -> None:
        from core.watch_catalog import rank_research_urls

        got = rank_research_urls(
            [
                {"url": "https://promo-codes.example/mcluck-bonus"},
                {"url": "https://www.mcluck.com/providers"},
                {"url": "https://casinoreview.example/mcluck"},
                {"url": "https://www.mcluck.com/terms"},   # legal pages never
                {"url": "not-a-url"},
            ],
            brand_host="mcluck.com",
        )
        assert got[0] == "https://www.mcluck.com/providers"
        assert got[-1] == "https://promo-codes.example/mcluck-bonus"
        assert all("terms" not in u for u in got)

    @pytest.mark.asyncio
    async def test_research_fills_what_the_site_missed_and_marks_the_source(
        self, wm, monkeypatch
    ) -> None:
        import core.watch_observe as wo
        from tools.watch import tools as T

        await wm.add_subject(company_id="c1", name="McLuck", url="https://www.mcluck.com")

        async def fake_collect(start_url, **kw):
            # the brand's own pages say nothing about providers
            return [{"url": "https://www.mcluck.com/providers", "title": "Providers",
                     "text": "Our lobby is powered by great games.", "error": None, "method": "http"}]

        async def fake_search(query, *, api_key, **kw):
            assert api_key == "sk-test"
            return [{"url": "https://review.example/mcluck", "title": "review", "snippet": ""}]

        async def fake_fetch(url, **kw):
            assert kw.get("proxy_url") is None, "research must not spend the metered exit"
            return ("McLuck carries Pragmatic Play and Hacksaw Gaming titles.", None, "http")

        monkeypatch.setattr(wo, "collect_pages", fake_collect)
        monkeypatch.setattr(wo, "search_web", fake_search)
        monkeypatch.setattr(wo, "fetch_page_best_effort", fake_fetch)

        t = T.WatchCatalogCollectTool()
        t._watch_manager, t._config, t._browser_manager = wm, None, None
        t._vault = {"search_sh_api_key": "sk-test"}
        t._router = _Router([{"name": "Pragmatic Play", "detail": "", "index": 0},
                             {"name": "Hacksaw Gaming", "detail": "", "index": 1}])
        res = await t.execute({"company_id": "c1", "kinds": ["provider"]})
        assert res.success, res.error
        kinds = res.data["brands"][0]["kinds"]["provider"]
        assert kinds["from"] == "public research" and kinds["found"] >= 2
        rows = await wm.list_catalog("c1")
        assert {r.name for r in rows} == {"Pragmatic Play", "Hacksaw Gaming"}
        assert all(r.source_type == "third_party" for r in rows)
        assert all(r.geo_state == "n/a" for r in rows)  # no geo claim, no proxy

    @pytest.mark.asyncio
    async def test_a_few_teaser_titles_are_not_a_lobby_when_min_items_says_so(
        self, wm, monkeypatch
    ) -> None:
        """High 5 / Pulsz Bingo, 2026-09-01: three titles on the public page
        counted as 'answered', so the lobby was never read. min_items sets
        what counts as answered; below it research runs and, if still
        short, the kind is marked for sign-in."""
        import core.watch_observe as wo
        from tools.watch import tools as T

        await wm.add_subject(company_id="c1", name="High 5", url="https://www.high5casino.com")
        searched: list[str] = []

        async def teaser(start_url, **kw):
            return [{"url": "https://www.high5casino.com/games", "title": "Games",
                     "text": "Play Green Machine, Golden Knight and Shake the Sky.", "error": None,
                     "method": "http"}]

        async def fake_search(query, *, api_key, **kw):
            searched.append(query)
            return [{"url": "https://review.example/high-5", "title": "review", "snippet": ""}]

        async def fake_fetch(url, **kw):
            return ("High 5 Casino games: Green Machine, Golden Knight, Shake the Sky, Jaguar Wild.",
                    None, "http")

        monkeypatch.setattr(wo, "collect_pages", teaser)
        monkeypatch.setattr(wo, "search_web", fake_search)
        monkeypatch.setattr(wo, "fetch_page_best_effort", fake_fetch)
        t = T.WatchCatalogCollectTool()
        t._watch_manager, t._config, t._browser_manager = wm, None, None
        t._vault = {"search_sh_api_key": "sk-test"}
        t._router = _Router([{"name": "Green Machine", "detail": "", "index": 0},
                             {"name": "Golden Knight", "detail": "", "index": 1},
                             {"name": "Shake the Sky", "detail": "", "index": 2}])

        res = await t.execute({"company_id": "c1", "kinds": ["game"], "sign_in_if_missing": True})
        k = res.data["brands"][0]["kinds"]["game"]
        assert k["found"] == 3 and k["from"] == "brand site" and not k.get("needs_sign_in")
        assert searched == []                                   # three items answered the default

        res = await t.execute({"company_id": "c1", "kinds": ["game"], "sign_in_if_missing": True,
                               "min_items": 10})
        k = res.data["brands"][0]["kinds"]["game"]
        assert searched                                          # below the bar: research ran
        assert k["from"] == "brand site + public research" and 3 < k["found"] < 10   # site + each research page
        assert k["needs_sign_in"] is True                        # still short: the lobby is next
        assert res.data["needs_sign_in"] == ["High 5:game"]
        assert len(await wm.list_catalog("c1")) == 3             # the same titles twice is one row each

    @pytest.mark.asyncio
    async def test_nothing_public_marks_the_kind_for_sign_in_only_when_asked(
        self, wm, monkeypatch
    ) -> None:
        import core.watch_observe as wo
        from tools.watch import tools as T

        await wm.add_subject(company_id="c1", name="McLuck", url="https://www.mcluck.com")

        async def empty_pages(start_url, **kw):
            return [{"url": "https://www.mcluck.com/store", "title": "Store",
                     "text": "Sign in to see your prices.", "error": None, "method": "http"}]

        async def no_hits(query, *, api_key, **kw):
            return []

        monkeypatch.setattr(wo, "collect_pages", empty_pages)
        monkeypatch.setattr(wo, "search_web", no_hits)

        t = T.WatchCatalogCollectTool()
        t._watch_manager, t._config, t._browser_manager = wm, None, None
        t._vault = {"search_sh_api_key": "sk-test"}
        t._router = _Router([])
        quiet = await t.execute({"company_id": "c1", "kinds": ["coin_package"]})
        assert quiet.data["needs_sign_in"] == []  # not asked, so not suggested
        loud = await t.execute({"company_id": "c1", "kinds": ["coin_package"],
                                "sign_in_if_missing": True})
        assert loud.data["needs_sign_in"] == ["McLuck:coin_package"]


class TestAttribution:
    """A page reached by searching a brand may be about ten other brands —
    or about somebody's own business (2026-08-27: a game studio's services
    page gave LuckyLand a provider it does not carry)."""

    def test_comparison_and_unrelated_pages_are_refused(self) -> None:
        from core.watch_catalog import research_page_ok
        from core.watch_voice import brand_aliases

        ll = brand_aliases("LuckyLand Slots", "https://www.luckylandslots.com")
        assert not research_page_ok(
            "https://www.wagertalk.com/sites-like/luckyland", "LuckyLand " * 9,
            "LuckyLand Slots", ll,
        )
        assert not research_page_ok(
            "https://www.juegostudio.com/game-development-services",
            "We build games for clients", "LuckyLand Slots", ll,
        )
        assert not research_page_ok("https://x.example/z", "A page about tractors", "Pulsz", ["Pulsz"])

    def test_brand_pages_and_reviews_are_accepted(self) -> None:
        from core.watch_catalog import research_page_ok
        from core.watch_voice import brand_aliases

        assert research_page_ok("https://www.mcluck.com/providers", "…", "McLuck", ["McLuck"])
        # the short name in the URL is enough — "…/reviews/modo/"
        assert research_page_ok(
            "https://time2play.com/casinos/reviews/modo/", "Modo review",
            "Modo Casino", brand_aliases("Modo Casino", "https://www.modo.us"),
        )
        # or the brand named repeatedly in the text
        assert research_page_ok(
            "https://x.example/y", "Pulsz is great. Pulsz pays. Pulsz has games.",
            "Pulsz", ["Pulsz"],
        )


class TestProvenanceIsVisible:
    def test_a_ladder_read_off_a_review_site_says_so(self, tmp_path) -> None:
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        catalog = {
            "items": 2, "label": "Raw inventory…",
            "totals": {"coin_package": 2}, "third_party_only": ["Pulsz coin_package"],
            "brands": [{
                "subject_id": "s", "name": "Pulsz", "is_self": True,
                "counts": {"coin_package": 2}, "providers": [], "games_sample": [],
                "promotions": [], "customer_states": ["logged_out"], "observed_at": "2026-08-27",
                "sources": {"coin_package": ["third_party"]},
                "packages": [
                    {"name": "$1.99", "price_usd": 1.99, "coins": "30,000 Gold Coins",
                     "detail": "", "url": "https://review.example", "source_type": "third_party"},
                    {"name": "$4.99", "price_usd": 4.99, "coins": "79,500 Gold Coins",
                     "detail": "", "url": "https://www.pulsz.com/store", "source_type": "site"},
                ],
            }],
        }
        card = {"rows": [], "dimensions": []}
        out = tmp_path / "p.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=catalog)
        texts = []
        for sl in Presentation(str(out)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
            texts.append("\n".join(parts))
        slide = next(t for t in texts if "Pulsz (us) – Coins / Promotions" in t)
        # the caveat is on the brand's own page, not implied
        assert "Packages as reported by public reviews, not read off the store" in slide
        assert "$1.99" in slide and "$4.99" in slide


class TestLadderHygiene:
    """A price ladder read off review-site prose picks up junk rungs and
    tautologies (2026-08-27 deck: 'Card Crush 30$ → like 12 coins', and
    'Crown Coins $1.99 → $1.99')."""

    @pytest.mark.asyncio
    async def test_a_rung_without_a_readable_price_is_dropped(self) -> None:
        page = "Packages start at 30$ and you get like 12 coins. The $9.99 tier gives 25 Mystery Coins."
        r = _Router([
            {"name": "30$", "coins": "like 12 coins", "detail": "", "index": 0},
            {"name": "$9.99", "coins": "25 Mystery Coins", "detail": "", "index": 1},
        ])
        got = await extract_catalog(r, kind="coin_package", brand="B", page_text=page)
        assert [g["name"] for g in got] == ["$9.99"] and got[0]["price_usd"] == 9.99

    @pytest.mark.asyncio
    async def test_a_grant_that_only_repeats_the_price_is_blanked(self) -> None:
        page = "The $1.99 package is available."
        r = _Router([{"name": "$1.99", "coins": "$1.99", "detail": "Min. package", "index": 0}])
        got = await extract_catalog(r, kind="coin_package", brand="B", page_text=page)
        assert got[0]["coins_text"] == "" and got[0]["price_usd"] == 1.99

    def test_the_brand_page_shows_the_detail_as_the_description(self, tmp_path) -> None:
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        catalog = {
            "items": 1, "label": "Raw inventory…", "totals": {"coin_package": 1, "promotion": 40},
            "third_party_only": ["Crown coin_package"],
            "brands": [{
                "subject_id": "s", "name": "Crown", "is_self": False,
                "counts": {"coin_package": 1, "promotion": 2}, "providers": [], "games_sample": [],
                "sources": {"coin_package": ["third_party"]},
                "promotions": [{"name": "August Deal", "detail": "", "image": "", "url": "u"}],
                "packages": [{"name": "$1.99", "price_usd": 1.99, "coins": "",
                              "detail": "Its coin split is not published", "url": "u",
                              "source_type": "third_party"}],
                "customer_states": ["logged_out"], "observed_at": "2026-08-27",
            }],
        }
        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=catalog)
        texts = []
        for sl in Presentation(str(out)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
            texts.append("\n".join(parts))
        pkg = next(t for t in texts if "Crown – Coins / Promotions" in t)
        assert "Its coin split is not published" in pkg    # the detail is the Description column
        assert "$1.99\n$1.99" not in pkg                   # never the price twice
        assert "2 promotions on record" in pkg             # the brand's total, not just what fits
        # the cross-brand ladder and the twelve-promotion list are gone: the per-brand pages carry the data
        assert not any("as reported by public sources" in t or "promotions on record – " in t for t in texts)


class TestPerBrandDetail:
    """The client's own reference pages: '<Brand> – Coins / Promotions'
    with Package · Gold coins · Sweeps coins and Promotion · Benefit · How to
    claim · Frequency; '<Brand> – Loyalty Club' with Tier · Qualification ·
    Reward; and Providers / Games. Tables, one fact per cell — not a
    paragraph, and not a summary."""

    def _catalog(self, n_games: int = 70):
        provs = [f"Studio {i}" for i in range(60)]
        games = [f"Game {i}" for i in range(n_games)]
        return {
            "items": 140, "label": "Raw inventory…",
            "totals": {"provider": 60, "coin_package": 3, "promotion": 2, "loyalty_tier": 3, "game": 70},
            "third_party_only": [],
            "brands": [{
                "subject_id": "s", "name": "Modo", "is_self": False,
                "counts": {"provider": 60, "coin_package": 3, "promotion": 2, "loyalty_tier": 3, "game": 70},
                "providers": provs, "games_sample": games[:12], "games_full": games,
                "sources": {"provider": ["site"]},
                "packages": [
                    {"name": "$1.99", "price_usd": 1.99, "coins": "4,000 GC", "detail": "",
                     "url": "u", "source_type": "site", "gold_coins": 4000.0, "sweeps_coins": None},
                    {"name": "$9.99", "price_usd": 9.99, "coins": "50,000 GC + 25 SC",
                     "detail": "First purchase offer", "url": "u", "source_type": "site",
                     "gold_coins": 50000.0, "sweeps_coins": 25.0},
                    {"name": "$49.99", "price_usd": 49.99, "coins": "100,000 GC + 51 SC",
                     "detail": "VIP Offer", "url": "u", "source_type": "site",
                     "gold_coins": 100000.0, "sweeps_coins": 51.0},
                ],
                "promotions": [], "promotions_full": [
                    {"name": "Daily Login Bonus", "detail": "", "image": "", "url": "u",
                     "benefit": "1500GC + 0.2SC, rising to 2500GC and 0.25SC after 3 days",
                     "how_to_claim": "Login daily to claim the bonus", "frequency": "Daily"},
                    {"name": "Refer a friend", "detail": "Level 1 at $100 purchase", "image": "",
                     "url": "u", "benefit": "Up to 200K GC + 100 SC per friend",
                     "how_to_claim": "Share a unique referral link", "frequency": "Ad hoc"},
                ],
                "tiers": [
                    {"name": "Iron", "qualification": "N/A", "reward": "0% Weekly Coin Boost", "url": "u"},
                    {"name": "Bronze", "qualification": "500,000 per month", "reward": "25% Weekly Coin Boost", "url": "u"},
                    {"name": "Black Diamond", "qualification": "12,500,000,000 / year",
                     "reward": "VIP Club access, 100% Weekly Coin Boost", "url": "u"},
                ],
                "customer_states": ["logged_out"], "observed_at": "2026-08-27",
            }],
        }

    def _texts(self, path):
        from pptx import Presentation

        prs = Presentation(str(path))
        out = []
        for sl in prs.slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
            out.append("\n".join(parts))
        W, H = prs.slide_width, prs.slide_height
        oob = [sh for sl in prs.slides for sh in sl.shapes
               if sh.left is not None and (sh.left < 0 or sh.top < 0
                                           or sh.left + sh.width > W + 18288
                                           or sh.top + sh.height > H + 18288)]
        assert not oob, "something runs off the page"
        return out

    def test_coins_and_promotions_page_matches_the_reference(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=self._catalog())
        page = next(t for t in self._texts(out) if "Modo – Coins / Promotions" in t)
        assert "Coin package | Gold coins | Sweeps coins | Description" in page
        assert "$1.99 | 4,000 | – | –" in page                          # no SC, no note → honest dashes
        assert "$9.99 | 50,000 | 25 | First purchase offer" in page      # the note is the Description
        assert "Promotion | Benefit | How to claim | Frequency" in page
        assert "Daily Login Bonus | 1500GC + 0.2SC" in page and "| Login daily to claim the bonus | Daily" in page

    def test_loyalty_and_library_pages(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=self._catalog())
        texts = self._texts(out)
        loyalty = next(t for t in texts if "Modo – Loyalty Club" in t)
        assert "Loyalty Club tier | Qualification | Reward" in loyalty
        assert "Bronze | 500,000 per month | 25% Weekly Coin Boost" in loyalty
        assert "Black Diamond | 12,500,000,000 / year | VIP Club access" in loyalty
        library = next(t for t in texts if "Modo – Providers / Games" in t)
        assert "Studio 0" in library and "Studio 59" in library     # all sixty, not a tick
        assert "Game 0" in library and "Game 69" in library         # all seventy fit on the page
        assert "more in the workbook" not in library                # nothing was cut, so no claim it was

    def test_a_library_too_long_for_the_page_says_where_the_rest_is(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=self._catalog(n_games=300))
        library = next(t for t in self._texts(out) if "Modo – Providers / Games" in t)
        assert "Game 109" in library and "Game 110" not in library
        assert "+190 more in the workbook" in library

    def test_rows_collected_before_the_columns_existed_still_fill_them(self, tmp_path) -> None:
        """The 293 promotions already on record carry no benefit / claim /
        frequency; the summary derives them from the row's own words."""
        from core.watch_catalog import summarize_catalog

        class Row:
            def __init__(self, kind, name, detail="", meta=None):
                self.kind, self.name, self.detail, self.meta = kind, name, detail, meta or {}
                self.subject_id, self.source_url, self.image_path = "s", "u", ""
                self.source_type, self.customer_state, self.observed_at = "site", "logged_out", "2026-08-27"
                self.price_usd, self.coins_text, self.sort_index = None, "", 0

        class Subj:
            subject_id, name, is_self = "s", "Modo", False

        cat = summarize_catalog([Row("promotion", "Daily login bonus", "Log in every 24 hours; 1,500 GC")],
                                [Subj()])
        promo = cat["brands"][0]["promotions_full"][0]
        assert promo["frequency"] == "Daily" and promo["how_to_claim"] == "Log in"
        assert promo["benefit"] == "Log in every 24 hours; 1,500 GC"

    def test_a_brand_without_tiers_gets_no_loyalty_page(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        cat = self._catalog()
        cat["brands"][0]["tiers"] = []
        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=cat)
        assert not any("Loyalty Club" in t for t in self._texts(out))


class TestStructuredFields:
    def test_frequency_is_read_off_the_offers_own_words(self) -> None:
        from core.watch_catalog import parse_frequency

        assert parse_frequency("Daily login bonus", "Login every 24 hours") == "Daily"
        assert parse_frequency("Happy Hour", "Takes place nightly from 9-10 PM") == "Daily"
        assert parse_frequency("Midweek Madness", "Play every Tuesday, Wednesday, and Thursday") == "Weekly"
        assert parse_frequency("My Stash Jackpot", "Claimable every Monday, but expires after seven days") == "Weekly"
        assert parse_frequency("Refer-a-friend", "600 GC + 20 SC") == "Per referral"
        assert parse_frequency("Cinco de Mayo Bundle", "May 5, 2026 to May 31, 2026") == "Limited time"
        # a dated welcome offer is limited-time, whatever its trigger
        assert parse_frequency("August Deal", "08/31/2026 expire; UPON SIGN UP") == "Limited time"
        assert parse_frequency("First Purchase Bonus", "50,000 GC for $9.99") == "One-time"
        assert parse_frequency("Signup Promo", "No expiration date") == "One-time"   # "expiration" is not "expires"
        assert parse_frequency("Social Media Rewards", "Ongoing wagering period") == "Ongoing"
        assert parse_frequency("VIP program", "Play more to level up") == ""

    def test_claim_route_is_read_off_the_offers_own_words(self) -> None:
        from core.watch_catalog import parse_claim

        assert parse_claim("AMOE", "send a handwritten request card") == "Mail-in request"
        assert parse_claim("Referral Program", "UPON SIGN UP, NO PURCHASE NECESSARY") == "Share referral link"
        assert parse_claim("Daily Email Competition", "3 SC just for opting in") == "Opt in"
        assert parse_claim("Signup Promo", "Claim on your first purchase") == "First purchase"
        assert parse_claim("Welcome bonus", "when I completed the sign-up process") == "Sign up"
        assert parse_claim("sign-up bonus", "Once your account is fully verified") == "Sign up + verify"
        assert parse_claim("Daily reload", "you just need to sign into your account") == "Log in"
        assert parse_claim("Welcome Email Bonus", "keep an eye on your inbox") == "Watch inbox"
        assert parse_claim("Tournaments", "available on selected slots") == "Play qualifying games"
        assert parse_claim("Cinco de Mayo Bundle", "Seasonal discounted coin package") == "Purchase"
        assert parse_claim("Weekly Prize Draws", "100 winners selected at random") == ""
        # codes: a real one is quoted, a word after "code" is not one
        assert parse_claim("No-deposit bonus", "The promo code 'COVERSBONUS' activates it") == "Use code COVERSBONUS"
        assert parse_claim("Welcome Bonus (Promo Code Required)", "Sign up and receive GC") == "Enter promo code"
        assert parse_claim("Daily reload", "You don't need a bonus code to claim, just log in") == "Log in"
        assert parse_claim("Bonus", "No bonus code required; opt in") == "Opt in"

    def test_promo_fields_prefer_the_model_and_fill_its_blanks(self) -> None:
        from core.watch_catalog import promo_fields

        read = promo_fields("Daily login bonus", "Login every 24 hours",
                            {"benefit": "1,500 GC", "how_to_claim": "Open the app", "frequency": "daily"})
        assert read == {"benefit": "1,500 GC", "how_to_claim": "Open the app", "frequency": "daily"}
        old = promo_fields("Daily login bonus", "Login every 24 hours; base reward 1,500 GC", {})
        assert old == {"benefit": "Login every 24 hours; base reward 1,500 GC",
                       "how_to_claim": "Log in", "frequency": "Daily"}
        # legal boilerplate is not a benefit when the title carries the grant
        titled = promo_fields("Get 1.5M CC + 75 FREE SC", "T&Cs Apply; 18+. Void where prohibited.", {})
        assert titled["benefit"] == "1,500,000 GC + 75 SC"
        assert promo_fields("Giveaways", "", {})["benefit"] == ""

    def test_parse_coins_reads_the_grant_as_printed(self) -> None:
        from core.watch_catalog import parse_coins

        assert parse_coins("800,000 GC 50 SC") == (800000.0, 50.0)
        assert parse_coins("120K Gold Coins + 60 SC FREE") == (120000.0, 60.0)
        assert parse_coins("1,500,000 Crown Coins, 75 SC") == (1500000.0, 75.0)
        assert parse_coins("40 SC + 800K GC + wheel spin") == (800000.0, 40.0)
        assert parse_coins("30,000 Gold Coins") == (30000.0, None)
        assert parse_coins("Get 1.5M CC + 75 FREE SC") == (1500000.0, 75.0)     # Crown Coins' own abbreviation
        assert parse_coins("up to 1750000 WC + 30 FREE SC") == (1750000.0, 30.0)  # WOW Coins'
        assert parse_coins("25 Mystery Coins plus 5 Battle Cards")[0] == 25.0   # not 25 million
        assert parse_coins("a great deal") == (None, None)

    @pytest.mark.asyncio
    async def test_promotions_and_tiers_come_back_structured(self) -> None:
        page = ("Daily Login Bonus: 1500GC + 0.2SC. Login daily to claim. "
                "Loyalty Club: Bronze needs 500,000 per month and gives a 25% Weekly Coin Boost.")
        r = _Router([{"name": "Daily Login Bonus", "benefit": "1500GC + 0.2SC",
                      "how_to_claim": "Login daily to claim", "frequency": "daily", "index": 0}])
        got = await extract_catalog(r, kind="promotion", brand="Modo", page_text=page)
        assert got[0]["meta"] == {"benefit": "1500GC + 0.2SC", "how_to_claim": "Login daily to claim",
                                  "frequency": "daily"}
        r2 = _Router([{"name": "Bronze", "qualification": "500,000 per month",
                       "reward": "25% Weekly Coin Boost", "index": 1}])
        got2 = await extract_catalog(r2, kind="loyalty_tier", brand="Modo", page_text=page)
        assert got2[0]["name"] == "Bronze" and got2[0]["sort_index"] == 1
        assert got2[0]["meta"]["reward"] == "25% Weekly Coin Boost"

    @pytest.mark.asyncio
    async def test_package_numbers_are_parsed_not_trusted(self) -> None:
        page = "The $9.99 pack grants 50,000 GC + 25 SC on your first purchase."
        r = _Router([{"name": "$9.99", "coins": "50,000 GC + 25 SC", "detail": "first purchase",
                      "gold_coins": 999, "sweeps_coins": 999, "index": 0}])   # model's numbers are wrong
        got = await extract_catalog(r, kind="coin_package", brand="B", page_text=page)
        assert got[0]["meta"] == {"gold_coins": 50000.0, "sweeps_coins": 25.0}

    @pytest.mark.asyncio
    async def test_meta_round_trips_through_the_register(self, wm) -> None:
        subj = await wm.add_subject(company_id="c1", name="Modo")
        await wm.add_catalog_item(company_id="c1", subject_id=subj.subject_id, brand_name="Modo",
                                  kind="loyalty_tier", name="Bronze", sort_index=1,
                                  meta={"qualification": "500,000", "reward": "25% boost"})
        await wm.add_catalog_item(company_id="c1", subject_id=subj.subject_id, brand_name="Modo",
                                  kind="coin_package", name="$9.99", price_usd=9.99,
                                  coins_text="50,000 GC + 25 SC", meta={"gold_coins": 50000.0, "sweeps_coins": 25.0})
        s = await wm.catalog_summary("c1")
        b = s["brands"][0]
        assert b["tiers"] == [{"name": "Bronze", "qualification": "500,000", "reward": "25% boost", "url": ""}]
        assert b["packages"][0]["gold_coins"] == 50000.0 and b["packages"][0]["sweeps_coins"] == 25.0

    def test_known_review_urls_are_predictable(self) -> None:
        from core.watch_catalog import brand_slug, known_review_urls, rank_research_urls

        assert brand_slug("Crown Coins Casino") == "crown-coins"
        assert known_review_urls("Crown Coins Casino") == [
            "https://igamingfuture.com/sweepstakes-casinos/reviews/crown-coins/"
        ]
        ranked = rank_research_urls([
            {"url": "https://random-review.example/crown-coins"},
            {"url": "https://igamingfuture.com/sweepstakes-casinos/reviews/crown-coins/"},
        ])
        assert ranked[0].startswith("https://igamingfuture.com")

    def test_workbook_carries_the_new_columns(self, tmp_path) -> None:
        from openpyxl import load_workbook

        from core.watch_xlsx import render_scorecard_xlsx

        rows = [
            {"brand": "Modo", "kind": "coin_package", "name": "$9.99", "price_usd": 9.99,
             "gold_coins": 50000.0, "sweeps_coins": 25.0, "coins": "50,000 GC + 25 SC", "detail": "",
             "source_type": "site", "url": "u", "session": "logged_out", "observed_at": "2026-08-27", "sort_index": 0},
            {"brand": "Modo", "kind": "promotion", "name": "Daily Login Bonus", "benefit": "1500GC + 0.2SC",
             "how_to_claim": "Login daily", "frequency": "daily", "detail": "", "image": "", "url": "u",
             "session": "logged_out", "observed_at": "2026-08-27", "sort_index": 0},
            {"brand": "Modo", "kind": "loyalty_tier", "name": "Bronze", "qualification": "500,000",
             "reward": "25% boost", "url": "u", "session": "logged_out", "observed_at": "2026-08-27", "sort_index": 1},
        ]
        path = render_scorecard_xlsx({"rows": [], "dimensions": []}, dimensions=[], evidence=[],
                                     staleness=[], path=tmp_path / "b.xlsx", catalog_rows=rows)
        wb = load_workbook(path)
        assert {"Coin packages", "Promotions", "Loyalty tiers"} <= set(wb.sheetnames)
        assert [c.value for c in wb["Coin packages"][3]][:4] == ["Brand", "Price USD", "Gold coins", "Sweeps coins"]
        assert [c.value for c in wb["Coin packages"][4]][1:4] == [9.99, 50000.0, 25.0]
        assert [c.value for c in wb["Promotions"][4]][1:5] == ["Daily Login Bonus", "1500GC + 0.2SC", "Login daily", "daily"]
        assert [c.value for c in wb["Loyalty tiers"][4]][1:4] == ["Bronze", "500,000", "25% boost"]


class TestRegister:
    @pytest.mark.asyncio
    async def test_recollection_updates_rather_than_duplicating(self, wm) -> None:
        subj = await wm.add_subject(company_id="c1", name="Crown", url="https://c.example")
        row, is_new = await wm.add_catalog_item(
            company_id="c1", subject_id=subj.subject_id, brand_name="Crown",
            kind="coin_package", name="$29.99", coins_text="GC 700", price_usd=29.99,
            sort_index=1, source_url="https://c.example/store",
        )
        assert is_new and row["catalog_id"]
        # the ladder changes: same package, new grant — one row, updated
        row2, is_new2 = await wm.add_catalog_item(
            company_id="c1", subject_id=subj.subject_id, brand_name="Crown",
            kind="coin_package", name="$29.99", coins_text="GC 900 + SC 60", price_usd=29.99,
            sort_index=1, source_url="https://c.example/store", customer_state="registered",
        )
        assert not is_new2
        rows = await wm.list_catalog("c1")
        assert len(rows) == 1 and rows[0].coins_text == "GC 900 + SC 60"
        assert rows[0].customer_state == "registered"
        # never scored, and the evidence register is untouched
        assert await wm.list_evidence("c1") == [] and await wm.list_scores("c1") == []
        assert dedupe_key("Crown", "coin_package", "$29.99") == dedupe_key("crown", "coin_package", " $29.99 ")

    @pytest.mark.asyncio
    async def test_vocabularies_are_enforced(self, wm) -> None:
        subj = await wm.add_subject(company_id="c1", name="Crown")
        with pytest.raises(ValueError):
            await wm.add_catalog_item(company_id="c1", subject_id=subj.subject_id,
                                      kind="jackpot", name="x")
        with pytest.raises(ValueError):
            await wm.add_catalog_item(company_id="c1", subject_id=subj.subject_id,
                                      kind="game", name="x", customer_state="whale")

    @pytest.mark.asyncio
    async def test_summary_reads_ladders_and_counts(self, wm) -> None:
        us = await wm.add_subject(company_id="c1", name="Us", is_self=True)
        cc = await wm.add_subject(company_id="c1", name="Crown")
        for name, price, idx in (("$49.99", 49.99, 2), ("$9.99", 9.99, 0), ("$29.99", 29.99, 1)):
            await wm.add_catalog_item(company_id="c1", subject_id=cc.subject_id, brand_name="Crown",
                                      kind="coin_package", name=name, price_usd=price, sort_index=idx)
        for prov in ("Pragmatic Play", "NetEnt"):
            await wm.add_catalog_item(company_id="c1", subject_id=cc.subject_id, brand_name="Crown",
                                      kind="provider", name=prov)
        await wm.add_catalog_item(company_id="c1", subject_id=us.subject_id, brand_name="Us",
                                  kind="game", name="Sweet Bonanza")
        s = await wm.catalog_summary("c1")
        assert s["items"] == 6 and s["totals"]["coin_package"] == 3
        assert s["brands"][0]["name"] == "Us"  # ours first
        crown = next(b for b in s["brands"] if b["name"] == "Crown")
        assert [p["price_usd"] for p in crown["packages"]] == [9.99, 29.99, 49.99]  # ladder in order
        assert crown["counts"]["provider"] == 2
        snap = await wm.get_snapshot(await wm.take_snapshot("c1"))
        assert snap["catalog"]["items"] == 6


class TestDeckAndWorkbook:
    def _catalog(self):
        return {
            "items": 8, "label": "Raw inventory as printed…",
            "totals": {"provider": 2, "coin_package": 2, "promotion": 2, "game": 2},
            "brands": [{
                "subject_id": "s", "name": "Crown", "is_self": False,
                "counts": {"provider": 2, "coin_package": 2, "promotion": 2, "game": 2},
                "providers": ["Pragmatic Play", "NetEnt"],
                "packages": [{"name": "$9.99", "price_usd": 9.99, "coins": "GC 100,000 + SC 10",
                              "detail": "starter", "url": "u"},
                             {"name": "$29.99", "price_usd": 29.99, "coins": "GC 700,000 + SC 55",
                              "detail": "", "url": "u"}],
                "promotions": [{"name": "150% Extra Coins", "detail": "first purchase, 24h",
                                "image": "", "url": "u"}],
                "games_sample": ["Sweet Bonanza", "Gates of Olympus"],
                "customer_states": ["registered"], "observed_at": "2026-08-27",
            }],
        }

    def test_appendix_slides_appear_only_with_rows(self, tmp_path) -> None:
        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [{"name": "Crown", "is_self": False, "rank": 1, "provisional": False,
                          "overall": {"normalized_pct": 60.0, "coverage_pct": 70.0}, "dimensions": {}}],
                "dimensions": []}
        a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
        n = factual_narrative(card, None, [], [])
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=a)
        render_executive_deck(card, diff=None, judged=[], summary=n, gaps=[], evidence_count=1, path=b,
                              catalog=self._catalog())
        # the Game portfolio page + the brand's Coins/Promotions page + its
        # Providers/Games page (no loyalty tiers in this fixture); the
        # cross-brand summaries are gone
        assert len(Presentation(str(b)).slides) == len(Presentation(str(a)).slides) + 3
        texts = []
        for sl in Presentation(str(b)).slides:
            parts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
            for sh in sl.shapes:
                if sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
            texts.append("\n".join(parts))
        assert any("Pragmatic Play" in t for t in texts)
        assert any("$29.99" in t and "GC 700,000 + SC 55" in t for t in texts)
        assert any("150% Extra Coins" in t for t in texts)
        assert any("Sweet Bonanza" in t for t in texts)

    def test_workbook_gains_one_sheet_per_kind(self, tmp_path) -> None:
        from openpyxl import load_workbook

        from core.watch_xlsx import render_scorecard_xlsx

        card = {"rows": [], "dimensions": []}
        rows = [
            {"brand": "Crown", "kind": "provider", "name": "NetEnt", "detail": "", "url": "u",
             "session": "logged_out", "observed_at": "2026-08-27", "sort_index": 0},
            {"brand": "Crown", "kind": "coin_package", "name": "$29.99", "price_usd": 29.99,
             "coins": "GC 700,000", "detail": "", "url": "u", "session": "registered",
             "observed_at": "2026-08-27", "sort_index": 1},
            {"brand": "Crown", "kind": "promotion", "name": "150% Extra", "detail": "24h",
             "image": "/shots/p.jpg", "url": "u", "session": "logged_out",
             "observed_at": "2026-08-27", "sort_index": 0},
            {"brand": "Crown", "kind": "game", "name": "Sweet Bonanza", "detail": "Pragmatic",
             "url": "u", "session": "logged_out", "observed_at": "2026-08-27", "sort_index": 0},
        ]
        plain = render_scorecard_xlsx(card, dimensions=[], evidence=[], staleness=[], path=tmp_path / "a.xlsx")
        assert "Providers" not in load_workbook(plain).sheetnames
        with_raw = render_scorecard_xlsx(card, dimensions=[], evidence=[], staleness=[],
                                         path=tmp_path / "b.xlsx", catalog_rows=rows)
        wb = load_workbook(with_raw)
        assert {"Providers", "Coin packages", "Promotions", "Games"} <= set(wb.sheetnames)
        assert wb["Coin packages"]["B4"].value == 29.99
        assert wb["Promotions"]["G4"].value == "/shots/p.jpg"   # Image column, after Benefit/How/Frequency/Terms


class TestGamePortfolio:
    """The client's own sheet: one row per studio, one column per brand.
    Their column A is the studio list; brands and reviews print the same
    studio five ways, so the matrix collapses spellings and keeps their
    order — and shows what is on their list that nothing has shown yet."""

    def test_spellings_collapse_to_one_studio(self) -> None:
        from core.watch_catalog import canonical_provider as c

        assert c("BGaming") == c("B Gaming") == c("BGAMING") == "bgaming"
        assert c("Relax Gaming") == c("Relax gaming") == c("Relax")
        assert c("2 By 2 Gaming") == c("2By2 Gaming") == c("2×2 Gaming")
        assert c("4TP (is this 4 the player?)") == c("4ThePlayer")
        assert c("Gamzik") == c("Gamzix") and c("BTG") == c("Big Time Gaming")
        assert c("Ela Games") == c("ElaGames") == c("Ela")
        assert c("Peter & Sons") == c("Peter and Sons") == c("Peter&Sons")
        assert c("VGW") == c("Virtual Gaming Worlds")
        assert c("Hacksaw Gaming") != c("Hacksaw RGS")           # the client lists both
        assert c("Gaming Corps") == "gamingcorps"                 # not a suffix when it leads
        assert c("Toucan Games") != c("Toucan Royale")

    def test_brand_labels_match_across_sheets(self) -> None:
        from core.watch_catalog import brand_key as k

        assert k("LuckyLand Casino") == k("LuckyLand Slots")
        assert k("High5 Casino") == k("High 5 Casino")
        assert k("Wow Vegas") == k("WOW Vegas")
        assert k("Pulsz") != k("Pulsz Bingo")

    def test_reads_the_clients_sheet(self, tmp_path) -> None:
        from core.watch_catalog import read_provider_universe

        p = tmp_path / "portfolio.csv"
        p.write_text(
            "﻿;;;\nGame Provider ;Pulsz ;Chumba Casino ;Spinfinite \n155;;;\n3 Oaks ;;;\n"
            "4TP (is this 4 the player?);;;\n;;;\nB Gaming ;;;\n",
            encoding="utf-8",
        )
        uni = read_provider_universe(p)
        assert uni["providers"] == ["3 Oaks", "4TP (is this 4 the player?)", "B Gaming"]
        assert uni["brands"] == ["Pulsz", "Chumba Casino", "Spinfinite"]

    def _items(self):
        return [
            {"brand": "Pulsz", "kind": "provider", "name": "BGaming", "detail": "", "source_type": "site"},
            {"brand": "Pulsz", "kind": "provider", "name": "B Gaming", "detail": "", "source_type": "third_party"},
            {"brand": "Chumba Casino", "kind": "provider", "name": "BGAMING", "detail": "", "source_type": "third_party"},
            {"brand": "Chumba Casino", "kind": "provider", "name": "Golden Rock Studios", "detail": "", "source_type": "site"},
            {"brand": "Pulsz", "kind": "game", "name": "Aztec Magic", "detail": "BGaming", "source_type": "site"},
            {"brand": "Pulsz", "kind": "game", "name": "Elvis Frog", "detail": "B Gaming", "source_type": "site"},
            {"brand": "Pulsz", "kind": "game", "name": "Slot X", "detail": "Jackpot Slots", "source_type": "site"},
        ]

    def test_matrix_follows_the_clients_list_and_marks_the_rest(self) -> None:
        from core.watch_catalog import provider_matrix

        brands = [{"name": "Chumba Casino", "is_self": False}, {"name": "Pulsz", "is_self": True}]
        m = provider_matrix(self._items(), brands, universe=["3 Oaks", "B Gaming"],
                            universe_brands=["Pulsz", "Chumba Casino", "Spinfinite"])
        names = [r["name"] for r in m["providers"]]
        assert names == ["3 Oaks", "BGaming", "Golden Rock Studios"]   # their order, then ours
        oaks, bg, gr = m["providers"]
        assert oaks["on_client_list"] and not oaks["observed"] and oaks["brand_count"] == 0
        assert bg["brands"]["Pulsz"] == {"carried": True, "source": "site", "games": 2}   # site beats review; spellings merge
        assert bg["brands"]["Chumba Casino"]["carried"] and bg["brands"]["Chumba Casino"]["source"] == "third_party"
        assert not gr["on_client_list"] and gr["brand_count"] == 1
        assert m["counts"] == {"observed": 2, "on_client_list": 2, "both": 1, "list_only": 1, "observed_only": 1}
        assert m["brands_only_on_client_list"] == ["Spinfinite"]
        assert m["brands_only_in_register"] == []
        assert "Jackpot Slots" not in names                                # a category is not a studio

    def test_without_a_list_the_matrix_is_what_was_observed(self) -> None:
        from core.watch_catalog import provider_matrix

        brands = [{"name": "Chumba Casino", "is_self": False}, {"name": "Pulsz", "is_self": True}]
        m = provider_matrix(self._items(), brands)
        assert [r["name"] for r in m["providers"]] == ["BGaming", "Golden Rock Studios"]   # most carried first
        assert all(r["on_client_list"] for r in m["providers"])                            # nothing to mark
        assert m["brands"] == ["Chumba Casino", "Pulsz"]

    def test_summary_carries_the_matrix(self) -> None:
        from core.watch_catalog import summarize_catalog

        class Row:
            def __init__(self, brand, kind, name, detail=""):
                self.subject_id, self.kind, self.name, self.detail = brand, kind, name, detail
                self.source_type, self.customer_state, self.observed_at = "site", "logged_out", "2026-09-01"
                self.price_usd, self.coins_text, self.sort_index, self.meta = None, "", 0, {}
                self.source_url, self.image_path = "u", ""

        class Subj:
            def __init__(self, sid, name, is_self=False):
                self.subject_id, self.name, self.is_self = sid, name, is_self

        cat = summarize_catalog([Row("p", "provider", "NetEnt"), Row("c", "provider", "NetEnt")],
                                [Subj("c", "Chumba"), Subj("p", "Pulsz", True)],
                                universe={"providers": ["NetEnt", "Zoot studios"], "brands": ["Pulsz"]})
        m = cat["matrix"]
        assert m["brands"] == ["Pulsz", "Chumba"]                      # ours first
        assert [r["name"] for r in m["providers"]] == ["NetEnt", "Zoot studios"]
        assert m["brands_only_in_register"] == ["Chumba"]
        assert summarize_catalog([], [Subj("c", "Chumba")])["matrix"] is None

    def test_workbook_gets_the_portfolio_sheet_first(self, tmp_path) -> None:
        from openpyxl import load_workbook

        from core.watch_xlsx import render_scorecard_xlsx

        rows = [{**it, "url": "u", "session": "logged_out", "observed_at": "2026-09-01", "sort_index": 0,
                 "is_self": it["brand"] == "Pulsz"} for it in self._items()]
        card = {"rows": [], "dimensions": []}
        path = render_scorecard_xlsx(card, dimensions=[], evidence=[], staleness=[], path=tmp_path / "w.xlsx",
                                     catalog_rows=rows,
                                     provider_universe={"providers": ["3 Oaks", "B Gaming"], "brands": ["Pulsz", "Spinfinite"]})
        wb = load_workbook(path)
        assert wb.sheetnames.index("Game portfolio") < wb.sheetnames.index("Providers")
        ws = wb["Game portfolio"]
        assert [c.value for c in ws[3]] == ["Game provider", "Pulsz", "Chumba Casino", "Brands", "On client list"]
        assert [c.value for c in ws[4]] == ["3 Oaks", None, None, 0, "yes"]
        assert [c.value for c in ws[5]] == ["BGaming", "● 2", "●", 2, "yes"]
        assert [c.value for c in ws[6]] == ["Golden Rock Studios *", None, "●", 1, "no"]
        assert any("Spinfinite" in str(c.value) for row in ws.iter_rows(min_row=7) for c in row if c.value)

    def test_a_bare_filename_resolves_inside_the_workspace(self, tmp_path, monkeypatch) -> None:
        from tools.watch.tools import _resolve_input

        class Cfg:
            workspace = str(tmp_path / "ws")

        (tmp_path / "ws" / "watch").mkdir(parents=True)
        (tmp_path / "ws" / "watch" / "game_portfolio.csv").write_text("Game Provider;Pulsz\n3 Oaks;\n")
        assert _resolve_input("game_portfolio.csv", Cfg()) == tmp_path / "ws" / "watch" / "game_portfolio.csv"
        monkeypatch.chdir(tmp_path)
        (tmp_path / "other.csv").write_text("Game Provider;Pulsz\nBGaming;\n")
        assert _resolve_input("other.csv", Cfg()) == tmp_path / "other.csv"          # the working directory last
        assert _resolve_input("/abs/none.csv", Cfg()) == Path("/abs/none.csv")       # absolute paths pass through

    def test_every_pack_tool_offers_the_clients_sheet(self) -> None:
        from tools.watch import tools as T

        for cls in (T.WatchExecutiveDeckTool, T.WatchBoardReportTool, T.WatchScorecardTool, T.WatchCatalogTool):
            assert "providers_from" in cls().input_schema["properties"], cls.__name__

    def test_deck_paginates_every_studio(self, tmp_path) -> None:
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        provs = [f"Studio {i:02d}" for i in range(60)]
        brands = ["Chumba Casino", "Pulsz"]
        items = [{"brand": b, "kind": "provider", "name": p, "detail": "", "source_type": "site"}
                 for b in brands for p in provs]
        from core.watch_catalog import provider_matrix
        matrix = provider_matrix(items, [{"name": "Pulsz", "is_self": True}, {"name": "Chumba Casino", "is_self": False}])
        catalog = {"items": 120, "label": "", "totals": {"provider": 120}, "third_party_only": [], "matrix": matrix,
                   "brands": [{"subject_id": "p", "name": "Pulsz", "is_self": True, "counts": {"provider": 60},
                               "providers": provs, "packages": [], "promotions": [], "games_sample": [], "sources": {}},
                              {"subject_id": "c", "name": "Chumba Casino", "is_self": False, "counts": {"provider": 60},
                               "providers": provs, "packages": [], "promotions": [], "games_sample": [], "sources": {}}]}
        card = {"rows": [], "dimensions": []}
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, catalog=catalog)
        prs = Presentation(str(out))
        pages = [sl for sl in prs.slides
                 if any(sh.has_text_frame and "Game portfolio" in sh.text_frame.text for sh in sl.shapes)]
        assert len(pages) == 3                                             # 60 studios, 26 a page
        titles = [next(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and "Game portfolio" in sh.text_frame.text)
                  for sl in pages]
        assert "(1 of 3)" in titles[0] and "(3 of 3)" in titles[-1]
        cells = [c.text for sl in pages for sh in sl.shapes if sh.has_table for r in sh.table.rows for c in r.cells]
        assert "Studio 00" in cells and "Studio 59" in cells
        hdr = [c.text for c in [sh for sh in pages[0].shapes if sh.has_table][0].table.rows[0].cells]
        assert hdr[1] == "Pulsz  (us)" and hdr[2] == "Chumba"              # ours first, "Casino" dropped
        H = prs.slide_height
        assert all(sh.top + sh.height <= H for sl in pages for sh in sl.shapes if sh.has_table)


class TestOffersPaginate:
    def test_fifteen_brands_of_offers_take_two_pages(self, tmp_path) -> None:
        """2026-09-02: the register reached 15 brands and the offers table ran
        past the slide. Nine a page, titled (i of k), nothing below the footnote."""
        from pptx import Presentation

        from core.watch_deck import factual_narrative, render_executive_deck

        card = {"rows": [], "dimensions": []}
        offers = [{"brand": f"Brand {i:02d}", "is_self": i == 0,
                   "welcome": "New players can sign up for a welcome bundle of 200% extra coins plus free SC " * 1,
                   "ongoing": "Daily login bonus, weekly tournaments and social giveaways every single day"}
                  for i in range(15)]
        out = tmp_path / "d.pptx"
        render_executive_deck(card, diff=None, judged=[], summary=factual_narrative(card, None, [], []),
                              gaps=[], evidence_count=1, path=out, offers=offers)
        prs = Presentation(str(out))
        pages = [sl for sl in prs.slides
                 if any(sh.has_text_frame and "The offers on the table" in sh.text_frame.text for sh in sl.shapes)]
        assert len(pages) == 2
        titles = [next(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and "offers on the table" in sh.text_frame.text)
                  for sl in pages]
        assert "(1 of 2)" in titles[0] and "(2 of 2)" in titles[1]
        tables = [sh for sl in pages for sh in sl.shapes if sh.has_table]
        assert [len(t.table.rows) for t in tables] == [10, 7]                # 9 + 6 brands, plus headers
        assert all(t.top + t.height <= 6.72 * 914400 for t in tables)        # above the footnote line
        cells = [c.text for t in tables for r in t.table.rows for c in r.cells]
        assert "Brand 00  (us)" in cells and "Brand 14" in cells


class TestSignedInRead:
    """A 'registered' read must go through the browser, where the session
    lives — an HTTP fetch sees the logged-out site (2026-09-01/02: two
    registered re-reads wrote nothing while the lobbies were live)."""

    class _BM:
        def __init__(self) -> None:
            self.calls: list[tuple[str, dict]] = []
            self.here = ""

        async def call_tool(self, name, params):
            self.calls.append((name, params))
            if name == "browser_navigate":
                self.here = params["url"]
                return {"success": True}
            if name == "browser_click_text":
                if params["text"] in ("Providers", "Get Coins"):
                    self.here = f"clicked:{params['text']}"
                    return {"success": True, "matchedText": params["text"]}
                raise RuntimeError("no such text")
            if name == "browser_get_html":
                body = {"clicked:Providers": "<ul><li>Pragmatic Play</li><li>Hacksaw Gaming</li></ul>",
                        "clicked:Get Coins": "<div>$4.99 79,500 GC + 5 SC</div>"}.get(
                    self.here, "<div class='grid'><h3>Money Train 2</h3><h3>Scarab Surge</h3></div>")
                return {"success": True, "html": body}
            if name == "browser_extract":
                return {"success": True, "text": ""}
            return {"success": True}

    @pytest.mark.asyncio
    async def test_reads_lobby_then_clicks_through_and_names_the_kinds(self, monkeypatch) -> None:
        import core.watch_observe as wo
        from core.watch_catalog import rank_catalog_pages, read_signed_in_pages

        async def no_consent(bm, **kw):
            return 0

        monkeypatch.setattr(wo, "dismiss_consent", no_consent)
        bm = self._BM()
        pages = await read_signed_in_pages(bm, "https://www.pulsz.com/", ["provider", "game", "coin_package"])
        assert [p["title"] for p in pages] == ["Lobby games", "Providers", "Store – Get Coins"]
        assert all(p["method"] == "browser_session" for p in pages)
        assert "Money Train 2" in pages[0]["text"] and "Hacksaw Gaming" in pages[1]["text"]
        assert "$4.99" in pages[2]["text"]
        by_kind = rank_catalog_pages(pages)
        assert [p["title"] for p in by_kind["game"]] == ["Lobby games"]
        assert [p["title"] for p in by_kind["provider"]] == ["Providers"]
        assert [p["title"] for p in by_kind["coin_package"]] == ["Store – Get Coins"]
        assert any(n == "browser_eval" and "scrollTo" in p["code"] for n, p in bm.calls)   # the grid lazy-loads
        assert not any(n == "browser_click_text" and p["text"] == "I AGREE" for n, p in bm.calls)  # nothing accepted

    @pytest.mark.asyncio
    async def test_a_registered_read_never_uses_http(self, wm, monkeypatch) -> None:
        import core.watch_observe as wo
        from tools.watch import tools as T

        await wm.add_subject(company_id="c1", name="Pulsz", url="https://www.pulsz.com")
        http_calls: list[str] = []
        session_calls: list[str] = []

        async def http_pages(start_url, **kw):
            http_calls.append(start_url)
            return []

        async def session_pages(bm, start_url, kinds, **kw):
            session_calls.append(start_url)
            return [{"url": start_url, "title": "Lobby games", "text": "Money Train 2 and Scarab Surge",
                     "error": None, "method": "browser_session"}]

        monkeypatch.setattr(wo, "collect_pages", http_pages)
        import core.watch_catalog as wc
        monkeypatch.setattr(wc, "read_signed_in_pages", session_pages)
        t = T.WatchCatalogCollectTool()
        t._watch_manager, t._config, t._browser_manager = wm, None, object()
        t._vault = {}
        t._router = _Router([{"name": "Money Train 2", "detail": "", "index": 0},
                             {"name": "Scarab Surge", "detail": "", "index": 1}])
        res = await t.execute({"company_id": "c1", "kinds": ["game"], "customer_state": "registered",
                               "research": False})
        assert res.success, res.error
        assert session_calls == ["https://www.pulsz.com"] and http_calls == []
        rows = await wm.list_catalog("c1")
        assert {r.name for r in rows} == {"Money Train 2", "Scarab Surge"}
        assert all(r.customer_state == "registered" and r.source_type == "site" for r in rows)


class TestTheAgentReadsTheLobby:
    """docs/90: the signed-in read is the agent clicking through the lobby,
    providers, store, promotions and VIP pages; every browser_get_html it
    makes is recorded with its URL and filed by the kind its report names."""

    @pytest.mark.asyncio
    async def test_recorded_pages_are_filed_by_the_agents_report(self) -> None:
        from types import SimpleNamespace

        from core.watch_catalog import rank_catalog_pages, read_signed_in_pages

        class BM:
            def __init__(self):
                self.here = "https://www.pulsz.com/"
                self.calls = []

            async def call_tool(self, name, params=None):
                self.calls.append(name)
                if name == "browser_eval":
                    import json as _j
                    return {"success": True, "resultJson": _j.dumps(self.here)}
                if name == "browser_get_html":
                    body = {"https://www.pulsz.com/": "<h3>Money Train 2</h3><h3>Scarab Surge</h3>",
                            "https://www.pulsz.com/store": "<div>$4.99 79,500 GC + 5 SC</div>"}[self.here]
                    if (params or {}).get("maxLength", 50000) < 1000:      # the agent's own short ask
                        return {"success": True, "html": body[:8]}
                    return {"success": True, "html": body}
                if name == "browser_navigate":
                    self.here = params["url"]
                return {"success": True}

        bm = BM()

        class Agent:
            _registry = SimpleNamespace(all_tools=lambda: [SimpleNamespace(name=n) for n in
                                                          ("browser_navigate", "browser_click", "browser_get_html",
                                                           "browser_scroll", "browser_eval", "vault_lookup")])
            excluded = None

            async def run_isolated(self, goal, *, excluded_tool_names=None, max_steps_override=None):
                Agent.excluded = set(excluded_tool_names or ())
                assert "you cannot navigate" in goal and "Do not accept" in goal
                await bm.call_tool("browser_get_html", {"maxLength": 200})  # the lobby, the agent's short ask
                bm.here = "https://www.pulsz.com/store"                    # the agent clicked "Get Coins"
                await bm.call_tool("browser_get_html", {"maxLength": 200})  # the store
                return SimpleNamespace(content="PAGE 1: lobby\nPAGE 2: store", steps_taken=6, tool_calls_made=[])

        pages = await read_signed_in_pages(bm, "https://www.pulsz.com/", ["game", "coin_package"], agent=Agent())
        assert [(p["title"], p["url"]) for p in pages] == [("Lobby games", "https://www.pulsz.com/"),
                                                            ("Store – Get Coins", "https://www.pulsz.com/store")]
        assert "Money Train 2" in pages[0]["text"] and "$4.99" in pages[1]["text"]   # the full DOM, not the stub
        assert all(p["method"] == "browser_session" and p["via"] == "agent" for p in pages)
        assert [p["kind"] for p in pages] == ["game", "coin_package"]                  # filed as the agent labelled
        by_kind = rank_catalog_pages(pages)
        assert by_kind["game"][0]["title"] == "Lobby games" and by_kind["coin_package"][0]["title"] == "Store – Get Coins"
        assert {"browser_navigate", "browser_eval", "vault_lookup"} <= Agent.excluded
        assert bm.call_tool is not None and not hasattr(bm.call_tool, "pages")   # the recorder is unwrapped after


    def test_a_labelled_page_is_filed_by_its_label_even_on_an_unfriendly_url(self) -> None:
        from core.watch_catalog import rank_catalog_pages

        pages = [{"url": "https://www.pulsz.com/sweepstakes-lobby", "title": "Lobby games", "text": "x",
                  "kind": "game"},
                 {"url": "https://www.pulsz.com/help", "title": "Promotions", "text": "y", "kind": "promotion"}]
        by_kind = rank_catalog_pages(pages)
        assert by_kind["game"][0]["title"] == "Lobby games" and by_kind["promotion"][0]["title"] == "Promotions"
